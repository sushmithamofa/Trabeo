r"""Turn the filled schedule workbook into the sizes.csv the Revit builder reads.

    python make_sizes_csv.py ..\standard\J1153_C_01_STR_2026-07-14.xlsx
    python make_sizes_csv.py BOOK.xlsx --out "..\..\New folder\1\sizes.csv"

Writes sizes.csv next to model.json. The workbook wins over the drawing for what
it fills in: sizes, a beam's top-of-beam offset, a slab's thickness and sunk depth.

Each row carries its Level. The builder only applies rows for the level it is
building (or rows with no level), so one workbook covering several floors can be
copied beside every floor's model.json unchanged.

Runs outside Revit, on normal Python with openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is needed:  python -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_OUT = os.path.join(os.path.dirname(HERE), "extracted", "sizes.csv")

# sheet -> which heading feeds which sizes.csv field. A missing entry means the
# sheet has no such value, so the drawing's own is kept.
SHEETS = {
    "Columns": {"mark": "Mark", "width": "Width (mm)", "depth": "Depth (mm)",
                "level": "Top Level"},
    "Beams": {"mark": "Mark", "width": "Width (mm)", "depth": "Depth (mm)",
              "level": "Level", "top": "Top of Beam Offset (mm)", "sunk": "Sunk Depth (mm)"},
    "Walls": {"mark": "Mark", "width": "Thickness (mm)", "level": "Top Level"},
    "Slabs": {"mark": "Mark", "depth": "Thickness (mm)", "level": "Level",
              "top": "Top Offset (mm)", "sunk": "Sunk Depth (mm)"},
}
FIELDS = ["mark", "width_mm", "depth_mm", "level", "top_mm", "sunk_mm", "from"]
HEADER_ROW = 4
FIRST_DATA_ROW = 6          # row 5 is the example row


def headers(ws):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW, column=c).value
        if v:
            out[str(v).replace("*", "").strip()] = c
    return out


def number(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", help="the filled schedule .xlsx")
    ap.add_argument("--out", default=DEFAULT_OUT, help="where to write sizes.csv")
    a = ap.parse_args(argv)

    if not os.path.exists(a.workbook):
        print("No such workbook: " + a.workbook, file=sys.stderr)
        return 1

    wb = load_workbook(a.workbook, data_only=True)
    rows, blank = [], []

    for sheet, spec in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        h = headers(ws)
        if spec["mark"] not in h:
            print(f"  {sheet}: no '{spec['mark']}' column, skipped")
            continue

        def cell(r, key):
            head = spec.get(key)
            return ws.cell(row=r, column=h[head]).value if head in h else None

        found = 0
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            mark = cell(r, "mark")
            if not mark or not str(mark).strip():
                continue
            mark = str(mark).strip()
            w, d = number(cell(r, "width")), number(cell(r, "depth"))
            top, sunk = number(cell(r, "top")), number(cell(r, "sunk"))
            # the two ways of saying a member drops: a negative offset or a sunk depth
            if sheet == "Slabs" and not sunk and top is not None and top < 0:
                sunk = -top
            if sheet == "Beams" and top is None and sunk:
                top = -sunk
            if sheet == "Slabs":
                top = None                 # the builder takes a slab's drop from sunk_mm
            if w is None and d is None and top is None and sunk is None:
                blank.append(f"{sheet}:{mark}")
                continue
            level = cell(r, "level")
            rows.append({"mark": mark, "width_mm": "" if w is None else w,
                         "depth_mm": "" if d is None else d,
                         "level": str(level).strip() if level else "",
                         "top_mm": "" if top is None else top,
                         "sunk_mm": "" if sunk is None else sunk, "from": sheet})
            found += 1
        print(f"  {sheet}: {found} marks with a value")

    if not rows:
        print("\nNothing to write: no mark in the workbook has a size yet. Fill the "
              "Width and Depth columns, or ask the consultant for them.")
        return 1

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", newline="", encoding="utf-8") as f:
        wr = csv.DictWriter(f, fieldnames=FIELDS)
        wr.writeheader()
        wr.writerows(rows)

    print(f"\nWrote {a.out} with {len(rows)} marks.")
    if blank:
        print(f"{len(blank)} marks in the workbook still have no size and were left out:")
        print("  " + ", ".join(blank[:20]) + (" ..." if len(blank) > 20 else ""))
        print("Those elements will be built at a default size.")
    if not rows:
        print("\nNothing was written. Fill the Width and Depth columns in the workbook "
              "first, or ask the consultant for them.")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
