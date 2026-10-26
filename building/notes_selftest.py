"""Notes for the consultant, the report, the reply, and carrying decisions to a revision.

    python -m building.notes_selftest
"""
from __future__ import annotations

import base64
import importlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from building import Building, notes                        # noqa: E402
from building.model import Element, Sheet, Value             # noqa: E402

try:
    from PIL import Image
except Exception:                                            # pictures need Pillow
    Image = None


def jpeg_url(w=60, h=40, colour=(200, 40, 40)) -> str:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(buf, "JPEG")
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def plan(b, file, layout="4-GROUND FLOOR#2", lv="GF", stem="DRG"):
    view = Sheet(id=f"{stem}:{layout}", file=file, title=f"{layout} PLAN", reader="sheets.registry", level=lv)
    b.add_sheet(view)
    s = Sheet(id=f"{stem}:{layout}/plan", file="x_view.dxf", reader="plan.legacy",
              extra={"view_sheet": view.id, "level_key": lv, "set": "JOB"})
    b.add_sheet(s)
    return s.id


def col(sid, n, mark, x, y):
    return Element(id=f"{sid}/column/{n}", type="column", sheet=sid, mark=mark,
                   geometry={"x": x, "y": y, "outline": [[x - 150, y - 150], [x + 150, y - 150],
                                                        [x + 150, y + 150], [x - 150, y + 150]]},
                   props={"mark": Value(mark)}, resolved={"mark": Value(mark), "width": Value(300), "depth": Value(300)})


def beam(sid, n, mark, a, c, w=230, d=600):
    return Element(id=f"{sid}/beam/{n}", type="beam", sheet=sid, mark=mark,
                   geometry={"start": list(a), "end": list(c), "mid": None},
                   props={"mark": Value(mark), "width": Value(w), "depth": Value(d)},
                   resolved={"mark": Value(mark), "width": Value(w), "depth": Value(d)})


def floor(file, shift=(0.0, 0.0), stem="DRG", revised=False):
    """One floor: four columns on a 6 m grid and beams between them. Revised: drawn
    elsewhere in model space (shift), B2 deeper, B4 gone, a new B5, element ids renumbered."""
    b = Building(project="JOB")
    sid = plan(b, file, stem=stem)
    sx, sy = shift
    pts = {"C1": (0, 0), "C2": (6000, 0), "C3": (6000, 6000), "C4": (0, 6000)}
    els = [col(sid, i, m, x + sx, y + sy) for i, (m, (x, y)) in enumerate(pts.items())]
    bs = [("B1", "C1", "C2", 600), ("B2", "C2", "C3", 750 if revised else 600), ("B3", "C3", "C4", 600)]
    if not revised:
        bs.append(("B4", "C4", "C1", 600))
    else:
        bs.append(("B5", "C1", "C3", 450))
    for i, (m, a, c, d) in enumerate(bs):
        (ax, ay), (cx, cy) = pts[a], pts[c]
        els.append(beam(sid, (i + 7) if revised else i, m, (ax + sx, ay + sy), (cx + sx, cy + sy), 230, d))
    for e in els:
        b.add(e)
    return b


def run_folder(project, name, b):
    d = os.path.join(project, "runs", name)
    os.makedirs(d)
    b.save(os.path.join(d, "building.json"))
    open(os.path.join(d, "review.html"), "w").write("<html></html>")
    return d


class Values(unittest.TestCase):
    def test_sizes_in_a_reply(self):
        self.assertEqual(notes.values_in("B47 = 230 X 450 mm"), [{"prop": "width", "value": 230}, {"prop": "depth", "value": 450}])
        self.assertEqual(notes.values_in("230x600"), [{"prop": "width", "value": 230}, {"prop": "depth", "value": 600}])
        self.assertEqual(notes.values_in("make it 450 deep"), [{"prop": "depth", "value": 450}])
        self.assertEqual(notes.values_in("thk 150", "slab"), [{"prop": "thickness", "value": 150}])
        self.assertEqual(notes.values_in("125mm thick", "slab"), [{"prop": "thickness", "value": 125}])
        self.assertEqual(notes.values_in("600 x 300", "column"), [{"prop": "width", "value": 300}, {"prop": "depth", "value": 600}])
        self.assertEqual(notes.values_in("see revised drawing"), [])
        self.assertEqual(notes.values_in(""), [])


@unittest.skipIf(Image is None, "Pillow is not installed")
class NotesAndReports(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="c2r_notes_")
        self.proj = os.path.join(self.tmp, "Hotel Main Block")
        os.makedirs(self.proj)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def add(self, **kw):
        note = {"el": "E1", "sheet": "S", "plan": "GROUND FLOOR", "lv": "GF", "type": "beam", "mark": "B47",
                "size": "230 x 600", "issue": "width: drawings disagree (230 / 300)", "at": [10, 20], **kw}
        return notes.change(self.proj, {"action": "add", "note": note, "images": [jpeg_url()]}, "me@x.com", "run1")

    def test_add_edit_close_delete(self):
        n = self.add(text="confirm width")
        self.assertEqual((n["id"], n["no"], n["state"], n["by"]), ("N1", 1, "open", "me@x.com"))
        self.assertTrue(notes.image_path(self.proj, n["images"][0]))
        n2 = notes.change(self.proj, {"action": "add", "note": {"type": "area", "at": [1, 2]}, "images": [],
                                      "text": ""} | {"note": {"type": "area", "at": [1, 2], "text": "lift core?"}}, "me")
        self.assertEqual(n2["id"], "N2")
        notes.change(self.proj, {"action": "edit", "id": "N1", "text": "confirm width please"})
        notes.change(self.proj, {"action": "screenshot", "id": "N1", "image": jpeg_url(30, 30)})
        notes.change(self.proj, {"action": "state", "id": "N1", "state": "closed"})
        d = notes.load(self.proj)
        self.assertEqual(d["notes"][0]["text"], "confirm width please")
        self.assertEqual(len(d["notes"][0]["images"]), 2)
        first, second = d["notes"][0]["images"]                                         # an edited screenshot
        notes.change(self.proj, {"action": "replace_image", "id": "N1", "image": first, "new": jpeg_url(40, 20)})
        now = notes.load(self.proj)["notes"][0]["images"]
        self.assertEqual((len(now), now[1]), (2, second))                              # in the old one's place
        self.assertNotEqual(now[0], first)
        self.assertFalse(notes.image_path(self.proj, first))                            # the old file gone
        self.assertTrue(notes.image_path(self.proj, now[0]))
        with self.assertRaises(ValueError):
            notes.change(self.proj, {"action": "replace_image", "id": "N1", "image": first, "new": jpeg_url()})
        d = notes.load(self.proj)
        self.assertEqual([n["id"] for n in notes.rows(self.proj)], ["N2"])              # closed: not sent
        self.assertEqual(len(notes.rows(self.proj, include_closed=True)), 2)
        imgs = list(d["notes"][0]["images"])
        notes.change(self.proj, {"action": "delete", "id": "N1"})
        self.assertFalse(any(notes.image_path(self.proj, i) for i in imgs))              # its pictures go too
        self.assertEqual(notes.load(self.proj)["next"], 3)                              # numbers never reused
        with self.assertRaises(ValueError):
            notes.change(self.proj, {"action": "add", "note": {}, "images": []})           # nothing to say
        with self.assertRaises(ValueError):
            notes.save_image(self.proj, 9, "data:text/html;base64,PGI+")                  # only pictures
        self.assertEqual(notes.image_path(self.proj, "..\\notes.json"), "")              # only its own files

    def test_excel_round_trip_with_the_consultants_answer(self):
        from openpyxl import load_workbook
        self.add(text="width 230 or 300?")
        notes.change(self.proj, {"action": "screenshot", "id": "N1", "image": jpeg_url()})   # two pictures
        self.add(text="depth?", mark="B12", el="E2")
        out = notes.write_xlsx(self.proj, notes.rows(self.proj), os.path.join(self.tmp, "q.xlsx"), "me@x.com")
        wb = load_workbook(out)
        ws = wb["Queries"]
        self.assertEqual(len(ws._images), 3)
        head = [c.value for c in ws[5]]
        self.assertEqual(head, notes.HEAD)
        # the consultant answers in the yellow column, and moves the second row up
        ref, rep = head.index("Ref") + 1, head.index("Consultant reply") + 1
        rows = {ws.cell(r, ref).value: r for r in range(6, ws.max_row + 1) if ws.cell(r, ref).value}
        ws.cell(rows["N1"], rep, "230 x 450, see rev R2")
        ws.cell(rows["N2"], rep, "")
        wb.save(out)
        got = notes.apply_reply(self.proj, out, "q.xlsx")
        self.assertEqual((got["answered"], got["unknown"]), (1, []))
        n1 = notes.load(self.proj)["notes"][0]
        self.assertEqual(n1["state"], "answered")
        self.assertEqual(n1["reply"]["values"], [{"prop": "width", "value": 230}, {"prop": "depth", "value": 450}])
        self.assertEqual(notes.load(self.proj)["notes"][1]["state"], "open")           # no answer yet

    def test_pdf_is_a_valid_file_with_its_pictures(self):
        self.add(text="width (230) or 300? \\ é – check")
        self.add(text="long " * 300, el="E2")
        out = notes.write_pdf(self.proj, notes.rows(self.proj), os.path.join(self.tmp, "q.pdf"), "me@x.com")
        raw = open(out, "rb").read()
        self.assertTrue(raw.startswith(b"%PDF-1.4"))
        self.assertTrue(raw.rstrip().endswith(b"%%EOF"))
        x = int(re.search(rb"startxref\n(\d+)", raw).group(1))
        self.assertTrue(raw[x:].startswith(b"xref"))
        offs = [int(o) for o in re.findall(rb"(\d{10}) 00000 n", raw)]
        for i, o in enumerate(offs, 1):                                                   # every object where it says
            self.assertTrue(raw[o:].startswith(b"%d 0 obj" % i), i)
        self.assertEqual(raw.count(b"/Subtype /Image"), 2)
        self.assertEqual(raw.count(b"/Type /Page "), 1)


class Revision(unittest.TestCase):
    def test_elements_found_in_a_shifted_revision(self):
        old, new = floor("A.dwg"), floor("A rev.dwg", shift=(25000.0, -4000.0), stem="DRG2", revised=True)
        o_sid, n_sid = "DRG:4-GROUND FLOOR#2/plan", "DRG2:4-GROUND FLOOR#2/plan"
        m = notes.match(old.of(sheet=o_sid), new.of(sheet=n_sid))
        mark = {e.id: e.mark for e in new.elements}
        got = {next(e.mark for e in old.elements if e.id == a): mark[b] for a, b in m.items()}
        self.assertEqual(got, {"C1": "C1", "C2": "C2", "C3": "C3", "C4": "C4", "B1": "B1", "B2": "B2", "B3": "B3"})
        pairs = notes.pair_plans(old, new, {"A.dwg": "A rev.dwg"})
        self.assertEqual(pairs, [(o_sid, n_sid)])
        ch = notes.changes(old, new, pairs)
        kinds = sorted((c["kind"], c["mark"]) for c in ch)
        self.assertEqual(kinds, [("added", "B5"), ("changed", "B2"), ("removed", "B4")])
        self.assertIn("depth 600 -> 750", next(c["text"] for c in ch if c["kind"] == "changed"))

    def test_curved_beams_and_numbers(self):
        """A curve is the same curve when its ends and middle are where they were (its bulge
        is not a distance); 300 and 300.0 are one size."""
        def curve(sid, n, dx=0.0):
            return Element(id=f"{sid}/beam/{n}", type="beam", sheet=sid, mark="B1",
                           geometry={"start": [0 + dx, 0], "end": [8000 + dx, 0], "mid": [4000 + dx, 2500]},
                           props={"mark": Value("B1")}, resolved={"mark": Value("B1"), "width": Value(300.0)})
        old = [curve("A", 0)]
        self.assertEqual(notes.match(old, [curve("B", 5)]), {"A/beam/0": "B/beam/5"})
        self.assertEqual(notes.match(old, [curve("B", 5, dx=900)]), {})                  # moved: not the same
        self.assertEqual(notes._diff({"width": 300.0, "mark": "b1"}, {"width": 300, "mark": "B1"}), [])
        self.assertEqual(notes._diff({"depth": 600}, {"depth": 750}), ["depth 600 -> 750"])

    def test_a_decision_is_not_a_change(self):
        """A revised drawing is compared drawing to drawing: a value decided in Review is not
        reported as changed by the revision."""
        from building.linker import link
        tmp = tempfile.mkdtemp(prefix="c2r_dec_")
        try:
            b = floor("A.dwg")
            p = os.path.join(tmp, "b.json")
            b.save(p)
            ov = {"elements": {"DRG:4-GROUND FLOOR#2/plan/beam/1": {"depth": 999}}, "prefixes": {}}
            decided = Building.load(p)
            link(decided, ov)
            e = next(x for x in decided.elements if x.id == "DRG:4-GROUND FLOOR#2/plan/beam/1")
            self.assertEqual(e.final("depth"), 999)
            drawn = notes.as_drawn(p, ov)
            e = next(x for x in drawn.elements if x.id == "DRG:4-GROUND FLOOR#2/plan/beam/1")
            self.assertEqual(e.final("depth"), 600)                                       # as drawn
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_decisions_and_notes_follow_into_the_revision(self):
        tmp = tempfile.mkdtemp(prefix="c2r_rev_")
        try:
            proj = os.path.join(tmp, "p")
            a, a2 = os.path.join(tmp, "A.dwg"), os.path.join(tmp, "A rev.dwg")
            open(a, "w").write("x")
            open(a2, "w").write("y")
            run_folder(proj, "2026-10-17_100000", floor(a))
            ov = {"project": {"levels": {"Level 0": 0}},
                  "elements": {"DRG:4-GROUND FLOOR#2/plan/beam/1": {"depth": 650},       # B2
                               "DRG:4-GROUND FLOOR#2/plan/beam/3": {"width": 300},       # B4: gone in the revision
                               "OTHER:plan/beam/1": {"depth": 1}},                        # not in either: untouched
                  "marks": {"B1@/GF": {"depth": 700}}, "exclude": ["DRG:4-GROUND FLOOR#2/plan/column/3"]}
            json.dump(ov, open(os.path.join(proj, "overrides.json"), "w"))
            os.makedirs(notes.folder(proj))
            notes.save(proj, {"notes": [{"id": "N1", "no": 1, "el": "DRG:4-GROUND FLOOR#2/plan/beam/1",
                                         "sheet": "DRG:4-GROUND FLOOR#2/plan", "text": "B2 depth?", "state": "sent",
                                         "images": []}], "next": 2})
            new = run_folder(proj, "2026-10-17_110000", floor(a2, (25000.0, -4000.0), "DRG2", revised=True))
            res = notes.after_analysis(proj, new, [a2], [{"old": a, "new": a2}], log=lambda m: None)
            self.assertEqual((res["moved"], res["relink"], res["notes_moved"]), (2, True, 1))
            got = json.load(open(os.path.join(proj, "overrides.json")))
            self.assertEqual(got["elements"], {"DRG2:4-GROUND FLOOR#2/plan/beam/8": {"depth": 650},
                                               "DRG:4-GROUND FLOOR#2/plan/beam/3": {"width": 300},
                                               "OTHER:plan/beam/1": {"depth": 1}})
            self.assertEqual(got["exclude"], ["DRG2:4-GROUND FLOOR#2/plan/column/3"])
            self.assertEqual(got["marks"], ov["marks"])                                   # mark decisions carry anyway
            self.assertEqual(got["project"], ov["project"])
            n = notes.load(proj)["notes"][0]
            self.assertEqual(n["el"], "DRG2:4-GROUND FLOOR#2/plan/beam/8")
            self.assertEqual(n["revision"]["result"], "changed")
            self.assertIn("depth 600 -> 750", n["revision"]["text"])
            self.assertEqual(len(notes.revision_changes(new)["changes"]), 3)
            self.assertIn("DRG:4-GROUND FLOOR#2/plan/beam/3", res["unmatched"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_an_unchanged_analysis_changes_nothing(self):
        """The same drawings analysed again: decisions and notes are left exactly as they were."""
        tmp = tempfile.mkdtemp(prefix="c2r_same_")
        try:
            proj = os.path.join(tmp, "p")
            a = os.path.join(tmp, "A.dwg")
            open(a, "w").write("x")
            run_folder(proj, "2026-10-17_100000", floor(a))
            ov = {"elements": {"DRG:4-GROUND FLOOR#2/plan/beam/1": {"depth": 650}}, "marks": {}, "exclude": []}
            ovp = os.path.join(proj, "overrides.json")
            json.dump(ov, open(ovp, "w"))
            before = open(ovp, "rb").read()
            notes.remember_drawings(proj, [a])
            new = run_folder(proj, "2026-10-17_110000", floor(a))
            res = notes.after_analysis(proj, new, [a], [], log=lambda m: None)
            self.assertEqual((res["moved"], res["relink"], res["changes"]), (0, False, 0))
            self.assertEqual(open(ovp, "rb").read(), before)
            self.assertIsNone(notes.revision_changes(new))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_a_revision_saved_over_the_old_file_is_noticed(self):
        tmp = tempfile.mkdtemp(prefix="c2r_over_")
        try:
            proj = os.path.join(tmp, "p")
            a = os.path.join(tmp, "A.dwg")
            open(a, "w").write("x")
            notes.remember_drawings(proj, [a])
            self.assertEqual(notes.revised_drawings(proj, [a], []), {})
            open(a, "w").write("a longer revision")
            self.assertEqual(notes.revised_drawings(proj, [a], []), {a: a})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class CadText(unittest.TestCase):
    def test_the_drawings_own_text_where_it_is_drawn(self):
        """Marks, sizes and grid names to lay over the plan: text, multiline text and block
        attributes (inside blocks too), never a hidden layer's, never outside the plan."""
        import ezdxf
        from building import cad_text
        tmp = tempfile.mkdtemp(prefix="c2r_text_")
        try:
            doc = ezdxf.new()
            msp = doc.modelspace()
            doc.layers.add("OFF").off()
            doc.layers.add("FROZEN").freeze()
            msp.add_text("B29a(230X600)", height=250, rotation=90, dxfattribs={"insert": (1000, 2000)})
            msp.add_mtext("S4\\P135 Thk", dxfattribs={"insert": (3000, 1000), "char_height": 200, "attachment_point": 5})
            blk = doc.blocks.new("GRID")
            blk.add_circle((0, 0), 400)
            blk.add_text("E", height=300, dxfattribs={"insert": (0, 0), "halign": 1, "valign": 2,
                                                    "align_point": (0, 0)})
            blk.add_attdef("NO", (0, -600), dxfattribs={"height": 150})
            ins = msp.add_blockref("GRID", (5000, 5000))
            ins.add_attrib("NO", "G7", (5000, 4400), dxfattribs={"height": 150})
            msp.add_text("hidden", height=200, dxfattribs={"insert": (100, 100), "layer": "OFF"})
            msp.add_text("frozen", height=200, dxfattribs={"insert": (100, 100), "layer": "FROZEN"})
            msp.add_text("far away", height=200, dxfattribs={"insert": (900000, 0)})
            msp.add_text("x" * 120, height=200, dxfattribs={"insert": (200, 200)})
            p = os.path.join(tmp, "plan.dxf")
            doc.saveas(p)
            got = {t[6]: t for t in cad_text.read(p, (-1000, -1000, 10000, 10000))}
            self.assertEqual(set(got), {"B29a(230X600)", "S4\n135 Thk", "E", "G7"})
            self.assertEqual(got["B29a(230X600)"][:4], [1000, 2000, 250.0, 90.0])
            self.assertEqual(got["S4\n135 Thk"][4:6], [1, 2])                      # centred both ways
            self.assertEqual(got["E"][:2] + got["E"][4:6], [5000, 5000, 1, 2])      # inside its bubble
            self.assertEqual(got["G7"][:2], [5000, 4400])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


@unittest.skipIf(Image is None, "Pillow is not installed")
class InTheApp(unittest.TestCase):
    def test_notes_through_the_app(self):
        import urllib.request
        tmp = tempfile.mkdtemp(prefix="c2r_napp_")
        os.environ["C2R_SETTINGS"] = os.path.join(tmp, "settings.json")
        import webapp
        webapp = importlib.reload(webapp)
        from http.server import ThreadingHTTPServer
        port = webapp.free_port(8960)
        srv = ThreadingHTTPServer(("127.0.0.1", port), webapp.H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{port}"

        def call(path, body=None, raw=None, ctype="application/json"):
            data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
            req = urllib.request.Request(base + path, data=data, method="POST" if data is not None else "GET",
                                         headers={"Content-Type": ctype})
            try:
                with urllib.request.urlopen(req) as r:
                    return r.status, r.read(), r.headers
            except urllib.error.HTTPError as exc:
                return exc.code, exc.read(), exc.headers
        try:
            proj = os.path.join(tmp, "Job")
            webapp.S.open(proj)
            run = os.path.join(proj, "runs", "2026-10-17_100000")
            os.makedirs(run)
            open(os.path.join(run, "review.html"), "w", encoding="utf-8").write("<html><body>R</body></html>")
            webapp.S.run_dir = run
            code, body, _ = call("/api/notes", {"action": "add", "note": {"el": "E1", "mark": "B47", "type": "beam",
                                                                          "text": "width?"}, "images": [jpeg_url()]})
            self.assertEqual(code, 200)
            n = json.loads(body)
            self.assertEqual(n["run"], "2026-10-17_100000")
            code, body, h = call("/notes/img/" + n["images"][0])
            self.assertEqual((code, h["Content-Type"]), (200, "image/jpeg"))
            self.assertEqual(call("/notes/img/..%5Coverrides.json")[0], 404)
            self.assertEqual(json.loads(call("/api/notes")[1])["notes"][0]["id"], "N1")
            # the review inside the app gets the notes; the file itself is unchanged
            page = call("/run/review.html")[1].decode()
            self.assertIn("Notes for the consultant", page)
            self.assertIn("/api/cadtext", page)
            self.assertEqual(json.loads(call("/api/cadtext?sheet=nope")[1])["t"], [])     # never an error page
            self.assertNotIn("Notes for the consultant", open(os.path.join(run, "review.html")).read())
            for fmt in ("xlsx", "pdf"):
                code, body, _ = call("/api/notes/report", {"fmt": fmt})
                self.assertEqual(code, 200, body)
                name = json.loads(body)["name"]
                self.assertTrue(name.startswith("Job - Queries - ") and name.endswith("." + fmt))
                code, data, h = call("/notes/report/" + urllib.request.quote(name))
                self.assertEqual(code, 200)
                self.assertIn("attachment", h["Content-Disposition"])
            self.assertEqual(json.loads(call("/api/notes")[1])["notes"][0]["state"], "sent")
            # the consultant's answer comes back in the Excel
            from openpyxl import load_workbook
            path = notes.folder(proj, "reports", notes.report_name(proj, "xlsx"))
            wb = load_workbook(path)
            ws = wb["Queries"]
            head = [c.value for c in ws[5]]
            ws.cell(6, head.index("Consultant reply") + 1, "300 x 600")
            buf = io.BytesIO()
            wb.save(buf)
            code, body, _ = call("/api/notes/reply?name=back.xlsx", raw=buf.getvalue(), ctype="application/octet-stream")
            self.assertEqual((code, json.loads(body)["answered"]), (200, 1))
            self.assertEqual(json.loads(call("/api/notes")[1])["notes"][0]["state"], "answered")
            self.assertEqual(call("/api/notes/reply?name=x.txt", raw=b"abc", ctype="application/octet-stream")[0], 400)
            # a revision in place of a drawing; the same file is refused
            a, a2 = os.path.join(tmp, "A.dwg"), os.path.join(tmp, "A rev.dwg")
            open(a, "w").write("x")
            open(a2, "w").write("y")
            webapp.S.files = [a]
            with self.assertRaises(RuntimeError):
                webapp._replace(0, a)
            webapp._replace(0, a2)
            self.assertEqual(webapp.S.files, [a2])
            v = webapp.S.view()
            self.assertEqual(v["files"][0]["rev_of"], "A.dwg")
            webapp.S.open(proj)                                                    # kept with the project
            self.assertEqual(webapp.S.revisions[0]["old"], a)
        finally:
            webapp.S.project, webapp.S.run_dir = None, None
            srv.shutdown()
            srv.server_close()
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
