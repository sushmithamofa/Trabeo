r"""Tests for the linker, the Excel and the review page (Stage 4).

    python -m building.linker_selftest

Each test rebuilds, small, a case met on the training drawings.
"""
from __future__ import annotations

import os
import tempfile
import unittest

from building import Building, Element, Evidence, Sheet, Value
from building.linker import link

V, P = "DWG:V1", "DWG:V1/plan"           # a registry view and the plan read from it


def building() -> Building:
    b = Building()
    b.add_sheet(Sheet(id=V, file="x.dwg", kind="framing_plan", level="L1", wing="B",
                      reader="sheets.registry"))
    b.add_sheet(Sheet(id=P, file="x.dwg", kind="framing_plan", level="Level 1", wing="B",
                      reader="plan.legacy", extra={"view_sheet": V, "level_key": "L1"}))
    return b


def ev(src="", sheet=P, text=""):
    return Evidence(reader="t", sheet=sheet, source=src, text=text)


def beam(b, n, start, end, mark="", width=230.0, depth=None, layer="BEAM"):
    props = {"width": Value(width, "read", 1.0, [ev("multiline")]), "kind": Value("beam", "read", 1.0, [ev()])}
    if depth:
        props["depth"] = Value(depth, "read", 1.0, [ev("drawing schedule")])
    return b.add(Element(id=f"{P}/beam/{n}", type="beam", sheet=P, level="Level 1", wing="B", mark=mark,
                         geometry={"start": list(start), "end": list(end)}, props=props,
                         extra={"_layer": layer, "_handle": f"h{n}"}))


def caption(b, n, at, mark, w=None, d=None, rot=0.0, family="beam"):
    props = {"family": Value(family, "read", 1.0, [ev("plan caption", V, mark)])}
    if w:
        props["width"] = Value(w, "read", 1.0, [ev("plan caption", V)])
        props["depth"] = Value(d, "read", 1.0, [ev("plan caption", V)])
    return b.add(Element(id=f"{V}/caption/{n}", type="caption", sheet=V, level="L1", wing="B", mark=mark,
                         geometry={"at": list(at), "text_height": 150.0, "rotation": rot},
                         props=props, extra={"raw": mark}))


def spec(b, n, mark, w, d, wing="B", level="L1", table="SCHEDULE OF BEAMS", part=""):
    return b.add(Element(id=f"T:{n}/spec", type="spec", sheet="T", level=level, wing=wing, mark=mark,
                         props={"family": Value("beam", "read", 1.0, [ev("table", "T")]),
                                "width": Value(w, "read", 1.0, [ev("table", "T")]),
                                "depth": Value(d, "read", 1.0, [ev("table", "T")])},
                         extra={"table": table, "part": part}))


class Captions(unittest.TestCase):
    def test_a_caption_goes_to_the_beam_it_runs_along(self):
        b = building()
        along = beam(b, 0, (0, 0), (6000, 0))                 # horizontal
        across = beam(b, 1, (3000, -3000), (3000, 3000))      # vertical, crossing near the caption
        caption(b, 0, (3150, 200), "B7", 230, 600, rot=0.0)
        link(b)
        self.assertEqual(along.final("mark"), "B7")
        self.assertIsNone(across.final("mark"))

    def test_a_caption_on_a_wide_girder_is_the_girders(self):
        """B72 was given GB7's caption: measured to the faces, not the centreline."""
        b = building()
        girder = beam(b, 0, (0, 0), (8000, 0), width=1500.0)
        side = beam(b, 1, (0, 900), (8000, 900), width=230.0)
        caption(b, 0, (4000, 550), "GB7", 1500, 850)          # on the girder, 550 from its centreline
        link(b)
        self.assertEqual(girder.final("mark"), "GB7")
        self.assertIsNone(side.final("mark"))

    def test_two_equally_close_beams_leave_the_caption_alone(self):
        b = building()
        beam(b, 0, (0, 0), (6000, 0), mark="B1")
        beam(b, 1, (0, 600), (6000, 600), mark="B2")
        caption(b, 0, (3000, 300), "B3", 230, 600)
        link(b)
        self.assertTrue(any(i.category == "caption" for i in b.issues))


class ColumnCaptions(unittest.TestCase):
    def column(self, b, n, x, y, layer="STR.COLUMN"):
        out = [[x - 115, y - 300], [x + 115, y - 300], [x + 115, y + 300], [x - 115, y + 300]]
        return b.add(Element(id=f"{P}/column/{n}", type="column", sheet=P, level="Level 1", wing="B",
                             geometry={"x": x, "y": y, "rotation": 90.0, "outline": out},
                             props={"width": Value(230.0, "read", 1.0, [ev()]),
                                    "depth": Value(600.0, "read", 1.0, [ev()])},
                             extra={"_layer": layer, "_handle": f"c{n}"}))

    def test_a_caption_naming_its_own_blocks_post_is_not_the_columns(self):
        """New block: 'SC1' from the steel-post block went to RCC columns 238, 239, 25A."""
        b = building()
        col = self.column(b, 0, 0.0, 0.0)
        cap = caption(b, 0, (640, 300), "SC1", family="column")
        cap.extra["block_outline"] = [0.0, 700.0, "Steel Col"]     # the post, 700 from the column
        link(b)
        self.assertFalse(col.final("mark"))
        self.assertTrue(any(i.category == "caption" and "Steel Col" in i.message for i in b.issues))

    def test_a_caption_beside_the_same_outline_still_attaches(self):
        """A plan drawn as one block: the column outline is in the block with its mark."""
        b = building()
        col = self.column(b, 0, 0.0, 0.0)
        cap = caption(b, 0, (400, 0), "C1", family="column")
        cap.extra["block_outline"] = [0.0, 0.0, "STR.COLUMN"]
        link(b)
        self.assertEqual(col.final("mark"), "C1")


class ShearWallsOnTheColumnLayer(unittest.TestCase):
    """Akanistha Winter Villa: SW1 ... SW18 drawn as rectangles on the column layer, each
    label beside its outline in one block - all were built as unmarked columns."""

    def col(self, b, n, x0, x1, t=300.0):
        out = [[x0, -t / 2], [x1, -t / 2], [x1, t / 2], [x0, t / 2]]
        return b.add(Element(id=f"{P}/column/{n}", type="column", sheet=P, level="Level 1", wing="B",
                             geometry={"x": (x0 + x1) / 2, "y": 0.0, "rotation": 0.0, "outline": out},
                             props={"kind": Value("column", "read", 1.0, [ev()]),
                                    "width": Value(t, "read", 1.0, [ev()]),
                                    "depth": Value(x1 - x0, "read", 1.0, [ev()])},
                             extra={"_layer": "COLUMN", "_handle": f"c{n}", "tag_distance": None, "repaired": False,
                                    "_keys": ["mark", "kind", "x", "y", "width", "depth", "rotation", "outline",
                                              "source_layer", "handle", "tag_distance", "repaired", "notes"]}))

    def test_an_outline_labelled_with_a_wall_mark_is_a_wall(self):
        from building.readers.legacy_plan import to_legacy
        b = building()
        sw = self.col(b, 0, -1000.0, 1000.0)                  # 300 x 2000, labelled SW3
        piece = self.col(b, 1, 1000.0, 2000.0)                # in line, touching: the rest of SW3
        squat = self.col(b, 2, 5000.0, 5600.0)                # 300 x 600 labelled SW4: a column
        cap = caption(b, 0, (0.0, 400.0), "SW3", family="wall")
        cap.extra["block_outline"] = [0.0, 0.0, "COLUMN"]
        cap2 = caption(b, 1, (5300.0, 400.0), "SW4", family="wall")
        cap2.extra["block_outline"] = [5300.0, 0.0, "COLUMN"]
        for _ in range(2):                                    # the linker runs again on its own result
            link(b)
            self.assertEqual((sw.type, piece.type, squat.type), ("wall", "wall", "column"))
            self.assertEqual((sw.final("mark"), piece.final("mark")), ("SW3", "SW3"))
            self.assertTrue(any(i.category == "wall" and "labelled SW3" in i.message for i in b.issues))
        walls = to_legacy(b, P)["walls"]
        self.assertEqual(len(walls), 2)
        self.assertEqual(sorted(walls[0]), sorted(["mark", "thickness", "start", "end", "outline", "source_layer",
                                                   "handle", "tag_distance", "notes"]))
        self.assertEqual(len(to_legacy(b, P)["columns"]), 1)

    def test_a_wall_label_written_beside_its_outline(self):
        """Akanistha re-saved 2026-10-24: the blocks holding each SW label with its rectangle
        exploded - the label lies beside the wall, not in it, with no outline of its own. The
        nearest outline, clearly nearer than the next, within five text heights, is the wall."""
        b = building()
        sw = self.col(b, 0, -1000.0, 1000.0)                  # 300 x 2000
        other = self.col(b, 1, 4000.0, 6000.0)                # another 300 x 2000, far off
        caption(b, 0, (0.0, 600.0), "SW3", family="wall")     # 450 above the wall's face
        link(b)
        self.assertEqual((sw.type, other.type), ("wall", "column"))
        self.assertEqual(sw.final("mark"), "SW3")

    def test_a_wall_label_between_two_outlines_names_neither(self):
        b = building()
        a = self.col(b, 0, -2200.0, -200.0)
        c = self.col(b, 1, 200.0, 2200.0)
        caption(b, 0, (0.0, 0.0), "SW3", family="wall")       # as near to both
        link(b)
        self.assertEqual((a.type, c.type), ("column", "column"))

    def test_an_outline_with_a_column_mark_stays_a_column(self):
        b = building()
        c = self.col(b, 0, -1000.0, 1000.0)
        cap = caption(b, 0, (0.0, 400.0), "SW3", family="wall")
        cap.extra["block_outline"] = [0.0, 0.0, "COLUMN"]
        caption(b, 1, (0.0, 0.0), "C1", family="column")
        link(b)
        self.assertEqual(c.type, "column")


class MarksFromAround(unittest.TestCase):
    """New block: a column-line beam marked once (NB5c) but cut into spans by every
    cross beam; an NS4 row written once and split into bays by SB beams."""

    def col(self, b, n, x, y):
        out = [[x - 150, y - 150], [x + 150, y - 150], [x + 150, y + 150], [x - 150, y + 150]]
        return b.add(Element(id=f"{P}/column/{n}", type="column", sheet=P, level="Level 1", wing="B",
                             geometry={"x": x, "y": y, "rotation": 0.0, "outline": out},
                             props={"width": Value(300.0, "read", 1.0, [ev()]), "depth": Value(300.0, "read", 1.0, [ev()])},
                             extra={"_layer": "COL", "_handle": f"c{n}"}))

    def slab(self, b, n, x0, y0, x1, y1, mark, written=True):
        return b.add(Element(id=f"{P}/slab/{n}", type="slab", sheet=P, level="Level 1", wing="B", mark=mark,
                             geometry={"outline": [[x0, y0], [x1, y0], [x1, y1], [x0, y1]]},
                             notes=[] if written else ["no slab mark in this bay: nearest mark used"],
                             extra={"_layer": "derived from beams"}))

    def test_a_span_past_a_crossing_beam_carries_the_beams_mark_up_to_a_column(self):
        b = building()
        self.col(b, 0, 0, 0), self.col(b, 1, 9000, 0)
        s1 = beam(b, 0, (0, 0), (3000, 0), mark="NB5c")
        s2 = beam(b, 1, (3000, 0), (6000, 0))
        s3 = beam(b, 2, (6000, 0), (9000, 0))
        beam(b, 3, (3000, 0), (3000, 3000), mark="B1"), beam(b, 4, (6000, 0), (6000, 3000), mark="B2")
        s4 = beam(b, 5, (9000, 0), (12000, 0))             # past the column: another beam
        link(b)
        self.assertEqual((s2.final("mark"), s3.final("mark")), ("NB5C", "NB5C"))
        self.assertFalse(s4.final("mark"))
        self.assertEqual(s1.final("mark"), "NB5C")

    def test_a_mark_written_at_both_ends_of_a_beam_line_names_the_spans_between(self):
        """NIFT: 4B7 written at both ends of a line of three spans over two columns."""
        b = building()
        for n, x in enumerate((0, 6000, 12000, 18000)):
            self.col(b, n, x, 0)
        s1 = beam(b, 0, (0, 0), (6000, 0))
        s2 = beam(b, 1, (6000, 0), (12000, 0))
        s3 = beam(b, 2, (12000, 0), (18000, 0))
        caption(b, 0, (1500, 200), "4B7")
        caption(b, 1, (16500, 200), "4B7")
        # a second line: different marks at its ends - nothing is taken
        t1 = beam(b, 3, (0, 8000), (6000, 8000))
        t2 = beam(b, 4, (6000, 8000), (12000, 8000))
        t3 = beam(b, 5, (12000, 8000), (18000, 8000))
        caption(b, 2, (1500, 8200), "4B1")
        caption(b, 3, (16500, 8200), "4B2")
        # a third: marked at one end only - nothing is taken
        u1 = beam(b, 6, (0, 16000), (6000, 16000))
        u2 = beam(b, 7, (6000, 16000), (12000, 16000))
        caption(b, 4, (1500, 16200), "4B3")
        link(b)
        self.assertEqual([s.final("mark") for s in (s1, s2, s3)], ["4B7", "4B7", "4B7"])
        self.assertEqual([s.final("mark") for s in (t1, t2, t3)], ["4B1", None, "4B2"])
        self.assertEqual([s.final("mark") for s in (u1, u2)], ["4B3", None])

    def test_a_bay_split_off_by_a_secondary_beam_takes_its_panels_mark(self):
        b = building()
        for n, (x, y) in enumerate(((0, 0), (6000, 0), (0, -1200), (6000, -1200))):
            self.col(b, n, x, y)
        beam(b, 0, (0, 0), (6000, 0), mark="NB3"), beam(b, 1, (0, -1200), (6000, -1200), mark="NB4")
        beam(b, 2, (3000, 0), (3000, -1200), mark="SB")     # ends on NB3 and NB4
        self.slab(b, 0, 115, -115, 2885, -1085, "NS4")
        bay = self.slab(b, 1, 3115, -115, 5885, -1085, "NS2", written=False)
        link(b)
        self.assertEqual(bay.resolved["mark"].value, "NS4")
        self.assertEqual(bay.resolved["mark"].status, "read")

    def test_across_a_beam_on_columns_the_nearest_mark_stays_a_working_value(self):
        """Stilt A wing: B11A, on a column at each end, between the sunk S1 strip and
        the unmarked bays above - S1 was carried up the whole middle row of bays."""
        b = building()
        for n, (x, y) in enumerate(((0, 0), (6000, 0), (0, -1200), (6000, -1200), (3000, 0), (3000, -1200))):
            self.col(b, n, x, y)
        beam(b, 0, (0, 0), (6000, 0), mark="NB3"), beam(b, 1, (0, -1200), (6000, -1200), mark="NB4")
        beam(b, 2, (3000, 0), (3000, -1200), mark="B11A")    # ends on two columns
        self.slab(b, 0, 115, -115, 2885, -1085, "S1")
        bay = self.slab(b, 1, 3115, -115, 5885, -1085, "S1", written=False)
        link(b)
        self.assertEqual(bay.resolved["mark"].status, "missing")

    def test_a_stair_slab_listed_in_the_slab_schedule_keeps_its_mark_and_thickness(self):
        """ALCAZAR 1st floor: ST1 written in the stair bay, 'ST-1,2,3 200 ... STAIR WAIST' in
        the SCHEDULE OF SLAB - had been set aside as 'a stair mark, not a slab's'."""
        b = building()
        b.add(Element(id="T:0/spec", type="spec", sheet="T", level="L1", wing="B", mark="ST1",
                      props={"family": Value("stair", "read", 1.0, [ev("table", "T")]),
                             "thickness": Value(200.0, "read", 1.0, [ev("table", "T")])},
                      extra={"table": "SCHEDULE OF SLAB:-", "part": ""}))
        s = self.slab(b, 0, 0, 0, 3000, 3000, "ST1")
        link(b)
        self.assertEqual((s.resolved["mark"].value, s.resolved["thickness"].value), ("ST1", 200.0))
        self.assertNotIn("tag_set_aside", s.extra)

    def test_across_a_main_beam_the_nearest_mark_stays_a_working_value(self):
        b = building()
        for n, (x, y) in enumerate(((0, 0), (6000, 0), (0, -2400), (6000, -2400))):
            self.col(b, n, x, y)
        beam(b, 0, (0, 0), (6000, 0), mark="NB3"), beam(b, 1, (0, -2400), (6000, -2400), mark="NB5")
        beam(b, 2, (0, -1200), (6000, -1200), mark="NB4")
        self.slab(b, 0, 115, -115, 5885, -1085, "NS4")
        bay = self.slab(b, 1, 115, -1315, 5885, -2285, "NS4", written=False)
        link(b)
        self.assertEqual(bay.resolved["mark"].status, "missing")


class Scope(unittest.TestCase):
    def zone(self, b, pts, layer="Main Block"):
        return b.add(Element(id=f"{V}/scope_zone/0", type="scope_zone", sheet=V,
                             geometry={"outline": [list(p) for p in pts]},
                             props={"kind": Value("other block", "read", 0.8, [ev("area", V)])},
                             extra={"layer": layer}))

    def test_what_lies_in_another_blocks_area_is_to_decide(self):
        """New block: beams B32a-e and the bays of the hatched 'Main Block' strip were
        modelled quietly as this block's."""
        b = building()
        inside = beam(b, 0, (0, 1000), (4000, 1000), mark="B32A")
        outside = [beam(b, n, (10000, n * 1000), (16000, n * 1000), mark=f"NB{n}") for n in range(1, 6)]
        self.zone(b, [(-500, -500), (5000, -500), (5000, 5000), (-500, 5000)])
        link(b)
        self.assertEqual(inside.resolved["scope"].status, "conflict")
        self.assertTrue(all("scope" not in e.resolved for e in outside))
        self.assertTrue(any(i.category == "scope" and "Main Block" in i.message for i in b.issues))
        link(b, {"elements": {inside.id: {"scope": "this block"}}})
        self.assertEqual(inside.resolved["scope"].status, "override")

    def test_an_area_round_nearly_everything_is_the_plans_own(self):
        b = building()
        els = [beam(b, n, (0, n * 1000), (4000, n * 1000)) for n in range(5)]
        self.zone(b, [(-500, -500), (5000, -500), (5000, 9000), (-500, 9000)])
        link(b)
        self.assertTrue(all("scope" not in e.resolved for e in els))


class Specs(unittest.TestCase):
    def test_a_spec_applies_by_mark_within_wing_and_level(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63")
        spec(b, 0, "B63", 230, 600)
        spec(b, 1, "B63", 230, 750, wing="C", table="WING C REVISED")
        spec(b, 2, "B63", 230, 750, part="club house", table="ROOF BEAMS CLUB HOUSE")
        spec(b, 3, "B63", 230, 900, level="L2")
        link(b)
        self.assertEqual((e.final("depth"), e.resolved["depth"].status), (600, "read"))

    def test_another_projects_mark_never_applies(self):
        """Sasan Gir's GB1 (450 deep) was 'contradicted' by Villa 5's GB1 (300)."""
        b = building()
        b.sheets[P].extra["set"] = "Sasangir"
        b.add_sheet(Sheet(id="T", file="villa.dwg", reader="sheets.registry", extra={"set": "Villa5"}))
        e = beam(b, 0, (0, 0), (6000, 0))
        caption(b, 0, (3000, 200), "GB1", 230, 450)
        spec(b, 0, "GB1", 230, 300, wing="", level="L1")
        link(b)
        self.assertEqual((e.final("depth"), e.resolved["depth"].status), (450, "read"))

    def test_table_and_old_extractor_disagree_is_a_conflict(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63", depth=750.0)
        spec(b, 0, "B63", 230, 600)
        link(b)
        self.assertEqual(e.resolved["depth"].status, "conflict")
        self.assertEqual(e.final("depth"), 600)              # the table is the stronger source
        self.assertTrue(any(i.category == "conflict" for i in b.issues))

    def test_the_old_readers_depth_from_another_wings_table_is_set_aside(self):
        """J1153: the old reader gave B-wing B63 the WING C REVISED row's 750."""
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63", depth=750.0)
        spec(b, 0, "B63", 230, 600)
        spec(b, 1, "B63", 230, 750, wing="C", table="WING C REVISED")
        link(b)
        self.assertEqual((e.final("depth"), e.resolved["depth"].status), (600, "read"))
        self.assertIn("another wing", e.extra["old_depth_set_aside"])

    def test_a_wing_note_on_the_plan_gives_its_beams_that_wings_sizes(self):
        """J1153 B-wing plan: 'WING C REVISED BEAMS' beside B21, B63 and BC10."""
        b = building()
        far = beam(b, 0, (0, 0), (6000, 0), mark="B21")
        b21 = beam(b, 1, (20000, 0), (26000, 0), mark="B21")
        b63 = beam(b, 2, (20000, 3000), (26000, 3000), mark="B63")
        caption(b, 0, (3000, 200), "B21")
        caption(b, 1, (23000, 200), "B21")
        caption(b, 2, (23000, 3200), "B63")
        b.sheets[V].extra["wing_notes"] = [{"wing": "C", "at": [23000, 1600], "text": "WING C REVISED BEAMS"}]
        for n, (mk, w) in enumerate((("B21", "B"), ("B63", "B"), ("B21", "C"), ("B63", "C"))):
            spec(b, n, mk, 230, 600 if w == "B" else 750, wing=w)
        link(b)
        self.assertEqual([x.final("depth") for x in (far, b21, b63)], [600, 750, 750])
        self.assertEqual((b21.wing, far.wing), ("C", "B"))

    def test_a_wing_note_whose_neighbours_are_not_its_schedule_changes_nothing(self):
        b = building()
        b21 = beam(b, 0, (20000, 0), (26000, 0), mark="B21")
        beam(b, 1, (20000, 3000), (26000, 3000), mark="B7")
        caption(b, 0, (23000, 200), "B21")
        caption(b, 1, (23000, 3200), "B7")
        b.sheets[V].extra["wing_notes"] = [{"wing": "C", "at": [23000, 2900], "text": "WING C REVISED BEAMS"}]
        spec(b, 0, "B21", 230, 600)
        spec(b, 1, "B21", 230, 750, wing="C")
        spec(b, 2, "B63", 230, 750, wing="C")
        link(b)
        self.assertEqual((b21.wing, b21.final("depth")), ("B", 600))   # B7 is nearer the note than B21
        self.assertTrue(any("keep this plan's wing" in i.message for i in b.issues))

    def test_a_missing_size_for_a_mark_written_in_a_schedule_row_is_called_a_reading_fault(self):
        """Any project: B66 written in a row the readers did not take ('B-66,83 230 600')
        must not be passed off as 'not given'."""
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B66")
        other = beam(b, 1, (0, 3000), (6000, 3000), mark="B175")
        b.sheets[V].extra["mentions"] = {"B66": [{"text": "B - 66,83 230 600", "sheet": V, "at": [0, -9000]}]}
        link(b)
        self.assertEqual(e.resolved["depth"].status, "missing")
        self.assertIn("66,83", e.extra["written_elsewhere"])
        self.assertTrue(any(i.category == "reading" and "B66" in i.message for i in b.issues))
        self.assertNotIn("written_elsewhere", other.extra)        # written nowhere: truly not given

    def test_a_layer_name_against_agreeing_captions_keeps_their_value_but_is_shown(self):
        """Changed 2026-10-10 (external review, point E): this was 'read', which hid
        that the drawing contradicts itself. The captions' value is still the one
        built; the element is now a conflict, so a person sees it and can pick."""
        b = building()
        b.add(Element(id=f"{V}/layer_size/0", type="layer_size", sheet=V,
                      props={"width": Value(230.0, "read", 0.9, [ev("layer name", V)]),
                             "depth": Value(600.0, "read", 0.9, [ev("layer name", V)])},
                      extra={"layer": "BEAM-230X600"}))
        e = beam(b, 0, (0, 0), (6000, 0), width=300.0, layer="BEAM-230X600")
        caption(b, 0, (3000, 200), "B3A", 300, 750)
        spec(b, 0, "B3A", 300, 750)
        link(b)
        self.assertEqual((e.final("depth"), e.resolved["depth"].status), (750, "conflict"))

    def layer(self, b, name, w, d, blocks=0):
        return b.add(Element(id=f"{V}/layer_size/{name}", type="layer_size", sheet=V,
                             props={"width": Value(w, "read", 0.9, [ev("layer name", V)]),
                                    "depth": Value(d, "read", 0.9, [ev("layer name", V)])},
                             extra={"layer": name, **({"blocks": blocks} if blocks else {})}))

    def test_a_layer_name_against_the_sections_is_shown_but_not_used(self):
        """New block NB3a: on BEAM-230X900, the beam sections say 230x750. 750 is
        kept; the 900 is no longer dropped without telling anyone."""
        b = building()
        self.layer(b, "BEAM-230X900", 230.0, 900.0)
        e = beam(b, 0, (0, 0), (6000, 0), layer="BEAM-230X900")
        caption(b, 0, (3000, 200), "NB3A")
        spec(b, 0, "NB3A", 230, 750)
        link(b)
        self.assertEqual((e.final("depth"), e.resolved["depth"].status), (750, "conflict"))
        self.assertTrue(any(i.category == "conflict" and "layer name" in i.message and "900" in i.message
                            for i in b.issues))

    def test_a_measured_width_against_the_layer_name_is_shown(self):
        """New block stair beam: faces 450 apart on BEAM-300X450."""
        b = building()
        self.layer(b, "BEAM-300X450", 300.0, 450.0)
        e = beam(b, 0, (0, 0), (6000, 0), width=450.0, layer="BEAM-300X450")
        link(b)
        self.assertEqual((e.final("width"), e.resolved["width"].status), (450, "conflict"))   # nothing guessed
        self.assertTrue(any(i.category == "conflict" and "plan geometry" in i.message and "layer name" in i.message
                            for i in b.issues))

    def test_a_depth_only_in_the_name_of_a_layer_holding_a_block_is_to_confirm(self):
        """New block B32b: 1325 deep only by BEAM-230X1325, which also holds the stair block."""
        b = building()
        self.layer(b, "BEAM-230X1325", 230.0, 1325.0, blocks=1)
        e = beam(b, 0, (0, 0), (6000, 0), layer="BEAM-230X1325")
        link(b)
        self.assertEqual(e.final("depth"), 1325)
        self.assertTrue(any(i.category == "confirm" and "BEAM-230X1325" in i.message for i in b.issues))

    def test_a_layer_name_that_agrees_says_nothing(self):
        b = building()
        self.layer(b, "BEAM-230X600", 230.0, 600.0)
        beam(b, 0, (0, 0), (6000, 0), layer="BEAM-230X600")
        caption(b, 0, (3000, 200), "B3")
        spec(b, 0, "B3", 230, 600)
        link(b)
        self.assertFalse([i for i in b.issues if i.category in ("conflict", "confirm")])

    def test_a_skipped_layer_named_as_structure_is_a_decision(self):
        """New block: 4 RCC projections on layer 'Projection', nothing built for them."""
        from building.model import Issue
        b = building()
        b.issues.append(Issue("info", "layers", "Layer 'Projection' carries 4 entities but is not in the layer "
                                                "map, so it was skipped. Add it to config.json if it holds structure.",
                              sheet=P))
        b.issues.append(Issue("info", "layers", "Layer 'A-FURN' carries 40 entities but is not in the layer "
                                                "map, so it was skipped.", sheet=P))
        link(b)
        got = [i for i in b.issues if i.category == "not modelled"]
        self.assertEqual(len(got), 1)
        self.assertIn("Projection", got[0].message)

    def test_nothing_known_is_missing_not_guessed(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B9")
        link(b)
        self.assertEqual(e.resolved["depth"].status, "missing")
        self.assertIsNone(e.final("depth"))

    def test_a_slab_tag_on_a_beam_is_set_aside(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="S4")
        caption(b, 0, (3000, 200), "B30A", 230, 600)
        link(b)
        self.assertEqual((e.final("mark"), e.resolved["mark"].status), ("B30A", "read"))
        self.assertIn("tag_set_aside", e.extra)


class Spans(unittest.TestCase):
    def test_one_beam_line_with_two_marks_splits_at_the_crossing_beam(self):
        b = building()
        line = beam(b, 0, (0, 0), (10000, 0))
        beam(b, 1, (4000, -3000), (4000, 3000), mark="B9")
        caption(b, 0, (2000, 200), "B3A", 300, 750)
        caption(b, 1, (7000, 200), "B3B", 230, 600)
        link(b)
        spans = sorted(b.of("beam_span"), key=lambda s: s.geometry["start"][0])
        self.assertEqual([(s.final("mark"), s.final("depth")) for s in spans], [("B3A", 750), ("B3B", 600)])
        self.assertAlmostEqual(spans[0].geometry["end"][0], 4000.0, delta=1.0)
        self.assertIn("split_into", line.extra)

    def test_no_support_between_captions_is_cut_midway_and_flagged(self):
        b = building()
        beam(b, 0, (0, 0), (10000, 0))
        caption(b, 0, (2000, 200), "B1", 230, 600)
        caption(b, 1, (6000, 200), "B2", 230, 450)
        link(b)
        self.assertTrue(any(i.category == "span split" for i in b.issues))

    def test_two_marks_written_on_the_same_spot_do_not_split_the_beam(self):
        """New block: 'NB2' and 'NB1' typed at one point on the right-hand NB2 (a slip)
        cut the beam midway and made half of it an 'NB1' of guessed size."""
        b = building()
        line = beam(b, 0, (0, 0), (6000, 0), mark="NB2")
        caption(b, 0, (3000, 200), "NB2")
        caption(b, 1, (3000, 200), "NB1")
        link(b)
        self.assertEqual(b.of("beam_span"), [])
        self.assertEqual(line.final("mark"), "NB2")
        self.assertTrue(any(i.category == "caption" and "NB1" in i.message and "same spot" in i.message
                            for i in b.issues), "the slip must be shown for review")

    def test_linking_twice_gives_the_same_spans(self):
        b = building()
        beam(b, 0, (0, 0), (10000, 0))
        beam(b, 1, (4000, -3000), (4000, 3000), mark="B9")
        caption(b, 0, (2000, 200), "B3A", 300, 750)
        caption(b, 1, (7000, 200), "B3B", 230, 600)
        link(b)
        link(b)
        self.assertEqual(len(b.of("beam_span")), 2)


class Overrides(unittest.TestCase):
    def test_decisions_win_and_keep_the_sources(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63", depth=750.0)
        f = beam(b, 1, (0, 3000), (6000, 3000), mark="B64")
        spec(b, 0, "B63", 230, 600)
        link(b, {"marks": {"B63@B/L1": {"depth": 600}}, "exclude": [f.id]})
        d = e.resolved["depth"]
        self.assertEqual((d.value, d.status), (600, "override"))
        self.assertIn("old extractor schedule: 750.0", [x.source for x in d.evidence])
        self.assertTrue(f.final("excluded"))

    def test_a_decision_for_another_level_does_not_apply(self):
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63")
        link(b, {"marks": {"B63@B/L2": {"depth": 900}}})
        self.assertNotEqual(e.final("depth"), 900)


class Outputs(unittest.TestCase):
    def test_excel_and_review_page_are_written_and_round_trip_decisions(self):
        from openpyxl import load_workbook
        from building import excel, review
        b = building()
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63", depth=750.0)
        spec(b, 0, "B63", 230, 600)
        link(b)
        with tempfile.TemporaryDirectory() as tmp:
            xl = excel.write(b, os.path.join(tmp, "r.xlsx"))
            wb = load_workbook(xl)
            self.assertIn("Beams", wb.sheetnames)
            ws = wb["Overrides"]
            ws.cell(row=5, column=1, value="B63@B/L1")
            ws.cell(row=5, column=2, value="depth")
            ws.cell(row=5, column=3, value=600)
            ws.cell(row=6, column=1, value="DWG:V1")
            ws.cell(row=6, column=2, value="level")
            ws.cell(row=6, column=3, value="L2")
            wb.save(xl)
            ov = excel.overrides_from_excel(xl)
            self.assertEqual(ov["marks"], {"B63@B/L1": {"depth": 600.0}})
            self.assertEqual(ov["sheets"], {"DWG:V1": {"level": "L2"}})
            html = review.write(b, os.path.join(tmp, "r.html"))
            text = open(html, encoding="utf-8").read()
            self.assertIn(e.id, text)
            self.assertNotIn("__DATA__", text)

    def test_masonry_is_drawn_left_out_not_as_structure_under_the_beams(self):
        """ALCAZAR 1st floor: the architect's brick walls inside the plan were drawn as
        elements with missing sizes, a band under every beam sitting on them - 'beams
        overlapping'. Never built, so drawn grey and dashed like everything left out."""
        from building import review
        b = building()
        beam(b, 0, (0, 0), (6000, 0), mark="B1", depth=450.0)
        ring = [[-115, -115], [6115, -115], [6115, 4115], [-115, 4115], [-115, -115],
                [115, 115], [115, 3885], [5885, 3885], [5885, 115], [115, 115]]
        for n, role in ((0, "masonry"), (1, "RC wall")):
            w = b.add(Element(id=f"{P}/wall/{n}", type="wall", sheet=P, level="Level 1", wing="B",
                              geometry={"outline": ring}, extra={"_layer": "BRICK WALL"}))
            w.resolved = {"role": Value(role, "read", 1.0, [ev()]), "width": Value(None, "missing", 0.0, [])}
        els = {e["id"]: e for s in review.data(b)["sheets"] for e in s["els"]}
        self.assertTrue(els[f"{P}/wall/0"]["out"])
        self.assertIn("masonry", els[f"{P}/wall/0"]["note"])
        self.assertFalse(els[f"{P}/wall/1"]["out"])                 # an RC wall is structure: drawn as read
        self.assertFalse(els[f"{P}/beam/0"]["out"])

    def test_every_section_and_detail_is_listed_with_what_was_taken_from_it(self):
        """'Sections are not used' (user, 2026-10-26): the review lists each section and detail
        - read (what it gave, how many elements use it), read with nothing usable found, or not
        read and why. A sheet's lone title is not listed."""
        from building import review
        b = building()
        b.add_sheet(Sheet(id="DWG:M2", file="x.dwg", kind="beam_details", title="B63 (230X600)",
                          reader="sheets.registry"))
        b.add_sheet(Sheet(id="DWG:M3", file="x.dwg", kind="section", title="SECTION A-A",
                          reader="sheets.registry", extra={"view": {"entities": 40}}))
        b.add_sheet(Sheet(id="DWG:S4", file="x.dwg", kind="stray", title="SECTION 1", reader="sheets.registry",
                          extra={"view": {"entities": 15}, "kind": {"reasons": ["only 15 entities"]}}))
        b.add_sheet(Sheet(id="DWG:S5", file="x.dwg", kind="stray", title="BEAM DETAILS", reader="sheets.registry",
                          extra={"view": {"entities": 1}}))
        b.add(Element(id="DWG:M2/spec", type="spec", sheet="DWG:M2", mark="B63",
                      props={"depth": Value(600.0, "read", 1.0, [ev("detail", "DWG:M2")])}))
        e = beam(b, 0, (0, 0), (6000, 0), mark="B63")
        e.props["depth"] = Value(600.0, "read", 1.0, [ev("detail", "DWG:M2")])
        got = {d["id"]: d for d in review.data(b)["details"]}
        self.assertEqual(sorted(got), ["DWG:M2", "DWG:M3", "DWG:S4"])
        self.assertEqual(got["DWG:M2"]["st"], "read")
        self.assertIn("a source for 1 element", got["DWG:M2"]["text"])
        self.assertEqual(got["DWG:M3"]["st"], "nothing")
        self.assertEqual(got["DWG:S4"]["st"], "not")
        self.assertIn("only 15 entities", got["DWG:S4"]["text"])

    def test_each_job_keeps_its_own_decisions_in_the_browser(self):
        """The browser kept decisions under the project name and the LENGTH of the sheet
        names, so two jobs could open each other's. Now: the job's folder and the sheets."""
        from building import review
        hotel = review._store(os.path.join("C:\\", "Jobs", "Hotel", "runs", "20261012_100000"))
        self.assertEqual(hotel, review._store(os.path.join("C:\\", "Jobs", "Hotel", "runs", "20261013_090000")))
        self.assertNotEqual(hotel, review._store(os.path.join("C:\\", "Jobs", "Mall", "runs", "20261012_100000")))
        self.assertEqual(review._store(os.path.join("C:\\", "out", "run1")),
                         os.path.normcase(os.path.join("C:\\", "out", "run1")))     # a run of its own
        self.assertEqual(review.data(building(), None, os.path.join("C:\\", "Jobs", "Hotel", "runs", "1"))["store"],
                         hotel)
        key = next(x for x in review.PAGE.splitlines() if x.startswith("const KEY="))
        self.assertIn("D.store", key)
        self.assertNotIn(".length", key)

    def test_the_cad_is_drawn_under_the_model_and_unused_outlines_stand_out(self):
        """External review, point G: 6 dropped New block columns were invisible in the
        review - only a line in a list. The drawing's own lines go under the model, and a
        column outline the model does not hold is highlighted."""
        import ezdxf
        from building import review
        b = building()
        cols = ColumnCaptions()
        cols.column(b, 0, 0.0, 0.0)                                          # modelled
        beam(b, 9, (0, 2000), (6000, 2000))                                  # the plan's extent
        with tempfile.TemporaryDirectory() as tmp:
            doc = ezdxf.new()
            msp = doc.modelspace()
            for x in (0.0, 6000.0):                                          # 6000: dropped
                msp.add_lwpolyline([(x - 115, -300), (x + 115, -300), (x + 115, 300), (x - 115, 300),
                                    (x - 115, -300)], dxfattribs={"layer": "STR.COLUMN"})
            msp.add_line((-1000, 2000), (7000, 2000), dxfattribs={"layer": "BEAM"})
            path = os.path.join(tmp, "plan.dxf")
            doc.saveas(path)
            b.sheets[P].dxf_file = path
            b.sheets[P].layer_map = {"STR.COLUMN": "column", "BEAM": "beam"}
            link(b)
            cad = next(s for s in review.data(b)["sheets"] if s["id"] == P)["cad"]
        self.assertEqual(len(cad["s"]) // 4, 9)                              # 4 + 4 sides, 1 line
        self.assertEqual([u["l"] for u in cad["u"]], ["STR.COLUMN"])
        self.assertGreater(min(p[0] for p in cad["u"][0]["p"]), 5000)

    def test_hidden_beams_are_shown(self):
        """A slab fold (the concealed / hidden beam at a sunk step, dashed on the SLAB
        layer) is built in Revit, so the review has to show it too: user saw them missing."""
        from building import review
        b = building()
        f = b.add(Element(id=f"{P}/fold/0", type="fold", sheet=P, level="Level 1", wing="B",
                          geometry={"start": [0, 0], "end": [4000, 0]},
                          props={"width": Value(200.0, "read", 1.0, [ev()]),
                                 "depth": Value(400.0, "read", 1.0, [ev()])}))
        link(b)
        els = [e for s in review.data(b)["sheets"] for e in s["els"]]
        got = next(e for e in els if e["id"] == f.id)
        self.assertEqual(got["t"], "hidden_beam")
        self.assertEqual(got["g"]["k"], "line")
        self.assertEqual(got["p"]["depth"]["v"], 400.0)


if __name__ == "__main__":
    unittest.main(verbosity=1)
