r"""Notes for the consultant: comments and screenshots on elements, sent as Excel or PDF,
answered back in the same Excel, and carried over to a revised drawing.

Everything lives in the project folder, apart from the model and the decisions:

    PROJECT\notes\notes.json          the notes (never read by the linker or the export)
    PROJECT\notes\img\N7_1.jpg        their screenshots
    PROJECT\notes\reports\...xlsx     the reports sent to the consultant
    PROJECT\notes\replies\...xlsx     the consultant's answers, as received
    PROJECT\notes\drawings_seen.json  each drawing's date at the last analysis (a revision
                                      saved over the old file is noticed by its date)

A note never changes the model. A value in the consultant's reply is only offered in
Review ("use consultant's value"); it becomes a decision when a person presses it.

After an analysis (after_analysis):
  * a decision or note whose element is no longer in the new analysis (a revised
    drawing gets new element ids) is moved to the same element in the new one - same
    type, same place (the revised plan may be shifted: the shift is measured from the
    columns both share), the same mark where both have one. Not found: left as it is
    and listed, never dropped;
  * for every revised drawing, what changed: elements added, removed, resized,
    renamed or moved (revision_changes.json in the run).

    python -m building.notes PROJECT report OUT.xlsx|OUT.pdf
"""
from __future__ import annotations

import base64
import datetime
import json
import math
import os
import re
import sys
import threading

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

DIR = "notes"
IMAGE_MB = 12                      # one screenshot, as sent by the page
STATES = ("open", "sent", "answered", "closed")
COMPARED = ("beam", "beam_span", "column", "slab", "wall", "footing", "raft")
SIZE_PROPS = ("width", "depth", "thickness", "diameter", "length")
_lock = threading.Lock()


# ---------------------------------------------------------------- storage
def folder(project: str, *sub: str) -> str:
    return os.path.join(project, DIR, *sub)


def load(project: str) -> dict:
    try:
        with open(folder(project, "notes.json"), encoding="utf-8") as f:
            d = json.load(f)
    except (OSError, ValueError):
        d = {}
    d.setdefault("notes", [])
    d.setdefault("next", 1 + max([n.get("no", 0) for n in d["notes"]] or [0]))
    return d


def save(project: str, d: dict) -> None:
    os.makedirs(folder(project), exist_ok=True)
    tmp = folder(project, "notes.json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, indent=1)
    os.replace(tmp, folder(project, "notes.json"))


def _now() -> str:
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M")


def save_image(project: str, note_no: int, data_url: str) -> str:
    """A screenshot sent by the page (a data: URL, JPEG or PNG) kept as a file; its name."""
    m = re.match(r"data:image/(jpeg|jpg|png);base64,(.+)$", data_url or "", re.S)
    if not m:
        raise ValueError("The screenshot did not arrive as a picture.")
    raw = base64.b64decode(m.group(2), validate=False)
    if len(raw) > IMAGE_MB * 1048576:
        raise ValueError("The screenshot is too large.")
    ext = "png" if m.group(1) == "png" else "jpg"
    if ext == "jpg" and not raw.startswith(b"\xff\xd8"):
        raise ValueError("The screenshot is not a JPEG picture.")
    if ext == "png" and not raw.startswith(b"\x89PNG"):
        raise ValueError("The screenshot is not a PNG picture.")
    os.makedirs(folder(project, "img"), exist_ok=True)
    k = 1
    while os.path.exists(folder(project, "img", f"N{note_no}_{k}.{ext}")):
        k += 1
    name = f"N{note_no}_{k}.{ext}"
    with open(folder(project, "img", name), "wb") as f:
        f.write(raw)
    return name


def image_path(project: str, name: str) -> str:
    """The file of a screenshot, or "" (only names this module made, in its own folder)."""
    name = os.path.basename(name or "")
    if not re.fullmatch(r"N\d+_\d+\.(jpg|png)", name):
        return ""
    p = folder(project, "img", name)
    return p if os.path.isfile(p) else ""


NOTE_FIELDS = ("el", "sheet", "plan", "lv", "type", "mark", "size", "issue", "status", "at",
               # made in Review's 3D view: its drawing set, the view to go back to, and - on the
               # architect's model - which model and object (building/arch_ui.py)
               "set", "view", "arch")
VIEW_BYTES = 8000                  # a note's kept view / architect's object: small by nature


def change(project: str, body: dict, who: str = "", run: str = "") -> dict:
    """One change from the Review page: add / edit / screenshot / state / used / delete."""
    act = body.get("action", "")
    with _lock:
        d = load(project)
        if act == "add":
            src = body.get("note") or {}
            text = str(src.get("text") or "").strip()
            imgs = body.get("images") or []
            if not text and not imgs:
                raise ValueError("Type a comment or take a screenshot first.")
            no = d["next"]
            n = {k: src.get(k) for k in NOTE_FIELDS if src.get(k) not in (None, "")}
            for k in ("view", "arch"):
                if k in n and (not isinstance(n[k], dict) or len(json.dumps(n[k])) > VIEW_BYTES):
                    del n[k]
            n.update({"id": f"N{no}", "no": no, "text": text, "images": [], "by": who, "date": _now(),
                      "state": "open", "run": run})
            for u in imgs:
                n["images"].append(save_image(project, no, u))
            d["notes"].append(n)
            d["next"] = no + 1
            save(project, d)
            return n
        n = next((x for x in d["notes"] if x.get("id") == body.get("id")), None)
        if n is None:
            raise ValueError("That note is no longer in the project.")
        if act == "edit":
            n["text"] = str(body.get("text") or "").strip()
            n["edited"] = _now()
        elif act == "screenshot":
            n["images"].append(save_image(project, n["no"], body.get("image") or ""))
        elif act == "replace_image":
            # a screenshot edited in Review (cropped, marked up): in the old one's place
            name = os.path.basename(body.get("image") or "")
            if name not in n["images"]:
                raise ValueError("That screenshot is no longer in the note.")
            new = save_image(project, n["no"], body.get("new") or "")
            n["images"][n["images"].index(name)] = new
            p = image_path(project, name)
            if p:
                os.remove(p)
        elif act == "drop_image":
            name = os.path.basename(body.get("image") or "")
            if name in n["images"]:
                n["images"].remove(name)
                p = image_path(project, name)
                if p:
                    os.remove(p)
        elif act == "state":
            st = body.get("state")
            if st not in STATES:
                raise ValueError("Unknown state.")
            n["state"] = st
        elif act == "used":
            n.setdefault("used", []).append({"prop": body.get("prop"), "value": body.get("value"),
                                             "date": _now(), "by": who})
        elif act == "delete":
            for name in n.get("images", []):
                p = image_path(project, name)
                if p:
                    os.remove(p)
            d["notes"].remove(n)
        else:
            raise ValueError("Unknown change.")
        save(project, d)
        return n


# ---------------------------------------------------------------- consultant's values
_NUM = r"(\d{2,4}(?:\.\d+)?)"


def values_in(text: str, type_: str = "beam") -> list[dict]:
    """Sizes written in a reply, as choices to offer: "230x450", "450 deep", "thk 150",
    "B47 = 230 X 600 mm". Never applied by itself."""
    t = " " + str(text or "").replace("×", "x").replace("X", "x").replace("*", "x") + " "
    out = []

    def add(p, v):
        v = float(v)
        v = int(v) if v == int(v) else v
        if 0 < v < 20000 and not any(o["prop"] == p for o in out):
            out.append({"prop": p, "value": v})
    if type_ == "slab":
        m = re.search(r"(?:thk|thick(?:ness)?|slab)\D{0,6}" + _NUM, t, re.I) or \
            re.search(_NUM + r"\s*(?:mm)?\s*(?:thk|thick)", t, re.I)
        if m:
            add("thickness", m.group(1))
        return out
    m = re.search(_NUM + r"\s*(?:mm)?\s*x\s*" + _NUM, t)
    if m:
        a, b = float(m.group(1)), float(m.group(2))
        if type_ == "column":
            add("width", min(a, b))
            add("depth", max(a, b))
        else:
            add("width", a)
            add("depth", b)
        return out
    m = re.search(r"(?:depth|deep|\bd\b)\D{0,4}" + _NUM, t, re.I) or re.search(_NUM + r"\s*(?:mm)?\s*(?:deep|depth)", t, re.I)
    if m:
        add("depth", m.group(1))
    m = re.search(r"(?:width|wide|\bb\b|\bw\b)\D{0,4}" + _NUM, t, re.I) or re.search(_NUM + r"\s*(?:mm)?\s*(?:wide|width)", t, re.I)
    if m:
        add("width", m.group(1))
    if type_ == "wall":
        m = re.search(r"(?:thk|thick(?:ness)?)\D{0,6}" + _NUM, t, re.I) or re.search(_NUM + r"\s*(?:mm)?\s*(?:thk|thick)", t, re.I)
        if m:
            add("width", m.group(1))
    return out


# ---------------------------------------------------------------- the report rows
def rows(project: str, include_closed: bool = False, extra: list | None = None) -> list[dict]:
    """The notes to send, in their order; extra: values still open with no note."""
    d = load(project)
    out = [n for n in d["notes"] if include_closed or n.get("state") != "closed"]
    out.sort(key=lambda n: n.get("no", 0))
    return out + list(extra or [])


def open_values(run_dir: str, noted: set) -> list[dict]:
    """Values still in conflict or missing, on elements with no note: rows without a
    screenshot, for "list every open value too"."""
    from building import Building
    from building.export import left_out
    from building.review import DRAWN, PROPS, SHOWN_AS, _key
    p = os.path.join(run_dir or "", "building.json")
    if not os.path.exists(p):
        return []
    b = Building.load(p)
    out = []
    for e in b.elements:
        if e.type not in DRAWN or e.id in noted or e.extra.get("split_into") or e.extra.get("merged_into"):
            continue
        s = b.sheets.get(e.sheet)
        if s is None or s.reader != "plan.legacy":
            continue
        if e.type == "stair" or left_out(e) or (e.type == "wall" and e.final("role") == "masonry"):
            continue
        if e.resolved.get("excluded"):
            continue
        bad = []
        for prop in PROPS[e.type]:
            k = _key(e.resolved.get(prop))
            if k in ("conflict", "missing"):
                bad.append(_issue_of(prop, e.resolved.get(prop), k))
        if not bad:
            continue
        out.append({"id": "", "no": 0, "el": e.id, "type": SHOWN_AS.get(e.type, e.type), "mark": e.final("mark") or "",
                    "lv": s.extra.get("level_key") or s.level or "", "plan": _plan_title(b, s),
                    "size": _size_text(e), "issue": "; ".join(bad), "text": "", "images": [], "auto": True})
    return out


def _issue_of(prop, v, k) -> str:
    if k == "missing":
        return f"{prop}: not given in the drawings"
    vals = []
    for ev in (v.evidence if v is not None else []):
        val = (ev.source or "").split(": ", 1)[-1] if ": " in (ev.source or "") else ""
        val = re.sub(r"^(\d+)\.0$", r"\1", val)
        if val and val not in vals:
            vals.append(val)
    return f"{prop}: drawings disagree" + (f" ({' / '.join(vals[:4])})" if vals else "")


def _plan_title(b, s) -> str:
    view = s.extra.get("view_sheet", "")
    return (b.sheets[view].title if view in b.sheets else "") or s.title or os.path.basename(s.file)


def _size_text(e) -> str:
    if e.type == "slab":
        t = e.final("thickness")
        return f"{_n(t)} thk" if t is not None else ""
    if e.final("diameter") is not None:
        return f"dia {_n(e.final('diameter'))}"
    w, dp = e.final("width"), e.final("depth")
    if w is None and dp is None:
        return ""
    return f"{_n(w) if w is not None else '?'} x {_n(dp) if dp is not None else '?'}"


def _n(v):
    try:
        f = float(v)
        return str(int(round(f))) if abs(f - round(f)) < 0.01 else f"{f:g}"
    except (TypeError, ValueError):
        return str(v)


def report_name(project: str, ext: str) -> str:
    job = re.sub(r'[<>:"/\\|?*]+', " ", os.path.basename(os.path.normpath(project))).strip() or "Project"
    return f"{job} - Queries - {datetime.date.today().isoformat()}.{ext}"


def mark_sent(project: str, ids: list[str]) -> None:
    with _lock:
        d = load(project)
        for n in d["notes"]:
            if n.get("id") in ids and n.get("state") == "open":
                n["state"] = "sent"
                n["sent"] = _now()
        save(project, d)


HEAD = ["No.", "Floor / plan", "Element", "Mark", "Size read", "What is unclear", "Screenshot", "Our comment",
        "Raised by / date", "Consultant reply", "Ref"]


def _floor(n) -> str:
    return " - ".join(x for x in (n.get("lv") or "", n.get("plan") or "") if x)


def _what(n) -> str:
    return (n.get("type") or "area").replace("_", " ")


# ---------------------------------------------------------------- Excel
def write_xlsx(project: str, items: list[dict], out: str, sender: str = "") -> str:
    """One row per note (a row per screenshot, the note's cells merged across them); the
    consultant writes in the yellow "Consultant reply" column and sends the file back."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    try:
        from openpyxl.drawing.image import Image as XLImage
        import PIL  # noqa: F401  (openpyxl needs it for pictures)
        pictures = True
    except Exception:
        pictures = False
    wb = Workbook()
    ws = wb.active
    ws.title = "Queries"
    job = os.path.basename(os.path.normpath(project))
    ws["A1"] = f"{job} - queries for the structural consultant"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Sent {datetime.date.today().strftime('%d %b %Y')}" + (f" by {sender}" if sender else "") + \
               f" - {len(items)} quer{'y' if len(items) == 1 else 'ies'}"
    ws["A3"] = "Please write your answer in the yellow \"Consultant reply\" column (sizes as 230 x 450, " \
               "or \"see revised drawing\") and send this file back."
    ws["A3"].font = Font(italic=True, color="555555")
    hr = 5
    widths = [6, 26, 11, 10, 12, 30, 72, 34, 18, 40, 8]
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, (h, w) in enumerate(zip(HEAD, widths), 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6FB2")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = box
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.column_dimensions[get_column_letter(len(HEAD))].hidden = True
    ws.freeze_panes = ws.cell(hr + 1, 1)
    yellow = PatternFill("solid", fgColor="FFF6C8")
    wrap = Alignment(vertical="top", wrap_text=True)
    r = hr + 1
    pic_w = 500                                   # px, about the Screenshot column
    for i, n in enumerate(items, 1):
        imgs = [image_path(project, x) for x in n.get("images") or []]
        imgs = [p for p in imgs if p] if pictures else []
        span = max(1, len(imgs))
        vals = [i, _floor(n), _what(n), n.get("mark") or "", n.get("size") or "", n.get("issue") or "",
                "" if imgs else ("(screenshot in the PDF)" if n.get("images") and not pictures else ""),
                n.get("text") or "", " ".join(x for x in (n.get("by") or "", n.get("date") or "") if x),
                (n.get("reply") or {}).get("text", "") if isinstance(n.get("reply"), dict) else "",
                n.get("id") or ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = wrap
            cell.border = box
            if HEAD[c - 1] == "Consultant reply":
                cell.fill = yellow
        for k, p in enumerate(imgs):
            try:
                pic = XLImage(p)
                scale = pic_w / float(pic.width or pic_w)
                pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
                ws.add_image(pic, f"G{r + k}")
                ws.row_dimensions[r + k].height = pic.height * 0.75 + 6
            except Exception:
                ws.cell(r + k, 7, "(screenshot could not be added)")
        if not imgs:
            lines = max(3, max(len(str(v)) // 30 + 1 for v in vals))
            ws.row_dimensions[r].height = min(15 * lines, 300)
        if span > 1:
            for c in range(1, len(HEAD) + 1):
                if c != 7:
                    ws.merge_cells(start_row=r, start_column=c, end_row=r + span - 1, end_column=c)
            for k in range(1, span):
                ws.cell(r + k, 7).border = box
        r += span
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    return out


def read_reply(path: str) -> dict:
    """{ref: reply text} from a returned report (any sheet with the Ref and Consultant
    reply columns; rows may have been moved; merged cells read from their first row)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    got = {}
    for ws in wb.worksheets:
        head_row, cols = None, {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
            names = {str(c.value).strip().lower(): c.column for c in row if c.value is not None}
            if "ref" in names and any(k.startswith("consultant reply") for k in names):
                head_row = row[0].row
                cols = {"ref": names["ref"],
                        "reply": next(v for k, v in names.items() if k.startswith("consultant reply"))}
                break
        if head_row is None:
            continue
        for r in range(head_row + 1, ws.max_row + 1):
            ref = ws.cell(r, cols["ref"]).value
            if not ref or not re.fullmatch(r"N\d+", str(ref).strip()):
                continue
            text = ws.cell(r, cols["reply"]).value
            text = "" if text is None else str(text).strip()
            if text:
                got[str(ref).strip()] = text
    return got


def apply_reply(project: str, path: str, original_name: str = "") -> dict:
    """The consultant's answers onto their notes. Returns what was matched."""
    got = read_reply(path)
    with _lock:
        d = load(project)
        by = {n["id"]: n for n in d["notes"]}
        matched, unknown = [], []
        for ref, text in got.items():
            n = by.get(ref)
            if n is None:
                unknown.append(ref)
                continue
            n["reply"] = {"text": text, "date": _now(), "file": original_name or os.path.basename(path),
                          "values": values_in(text, n.get("type") or "beam")}
            if n.get("state") != "closed":
                n["state"] = "answered"
            matched.append(ref)
        save(project, d)
    return {"answered": len(matched), "unknown": unknown, "rows": len(got)}


# ---------------------------------------------------------------- PDF (no library)
# Helvetica's widths (1/1000 em) for the printable ASCII characters
_HW = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278] + [556] * 10 + \
      [278, 278, 584, 584, 584, 556, 1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833,
       722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333, 556,
       556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556,
       500, 722, 500, 500, 500, 334, 260, 334, 584]


def _tw(s: str, size: float, bold=False) -> float:
    w = sum(_HW[ord(c) - 32] if 32 <= ord(c) < 127 else 556 for c in s)
    return w * size / 1000.0 * (1.06 if bold else 1.0)


def _wrap(text: str, width: float, size: float, bold=False) -> list[str]:
    out = []
    for para in str(text or "").replace("\r", "").split("\n"):
        line = ""
        for word in para.split(" "):
            t = (line + " " + word) if line else word
            if _tw(t, size, bold) <= width:
                line = t
                continue
            if line:
                out.append(line)
            while _tw(word, size, bold) > width and len(word) > 1:      # a very long word
                k = len(word)
                while k > 1 and _tw(word[:k], size, bold) > width:
                    k -= 1
                out.append(word[:k])
                word = word[k:]
            line = word
        out.append(line)
    return out


def _pdf_str(s: str) -> bytes:
    b = str(s).encode("cp1252", "replace")
    return b"(" + b.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def _jpeg_info(raw: bytes):
    """(width, height, components) of a JPEG, or None."""
    i = 2
    while i + 9 < len(raw):
        if raw[i] != 0xFF:
            i += 1
            continue
        m = raw[i + 1]
        if m in (0xD8, 0x01) or 0xD0 <= m <= 0xD7:
            i += 2
            continue
        seg = int.from_bytes(raw[i + 2:i + 4], "big")
        if 0xC0 <= m <= 0xCF and m not in (0xC4, 0xC8, 0xCC):
            h = int.from_bytes(raw[i + 5:i + 7], "big")
            w = int.from_bytes(raw[i + 7:i + 9], "big")
            return w, h, raw[i + 9]
        i += 2 + seg
    return None


def _as_jpeg(path: str):
    """A screenshot as JPEG bytes (PNG converted when Pillow is there), or None."""
    with open(path, "rb") as f:
        raw = f.read()
    if raw.startswith(b"\xff\xd8"):
        return raw
    try:
        import io
        from PIL import Image
        im = Image.open(io.BytesIO(raw)).convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88)
        return buf.getvalue()
    except Exception:
        return None


class _Pdf:
    W, H = 842.0, 595.0                       # A4 landscape, points

    def __init__(self):
        self.objs: list[bytes] = []
        self.pages: list[int] = []
        self.images: list[tuple[int, str]] = []
        self.ops: list[bytes] = []
        self.page_images: list[str] = []

    def obj(self, body: bytes) -> int:
        self.objs.append(body)
        return len(self.objs)

    def text(self, x, y, s, size=10, bold=False, grey=False):
        col = b"0.37 0.42 0.47 rg " if grey else b"0.09 0.13 0.17 rg "
        self.ops.append(col + b"BT /" + (b"F2" if bold else b"F1") + b" %.1f Tf %.1f %.1f Td " % (size, x, y) +
                        _pdf_str(s) + b" Tj ET")

    def rect(self, x, y, w, h, fill=None, stroke=(0.8, 0.8, 0.8)):
        s = b""
        if fill:
            s += b"%.3f %.3f %.3f rg %.1f %.1f %.1f %.1f re f " % (*fill, x, y, w, h)
        if stroke:
            s += b"%.3f %.3f %.3f RG 0.6 w %.1f %.1f %.1f %.1f re S" % (*stroke, x, y, w, h)
        self.ops.append(s)

    def image(self, jpeg: bytes, x, y, w, h):
        info = _jpeg_info(jpeg)
        if not info:
            return False
        iw, ih, comp = info
        cs = {1: b"/DeviceGray", 4: b"/DeviceCMYK"}.get(comp, b"/DeviceRGB")
        num = self.obj(b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace %s "
                       b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % (iw, ih, cs, len(jpeg))
                       + jpeg + b"\nendstream")
        k = min(w / iw, h / ih)
        dw, dh = iw * k, ih * k
        name = f"Im{num}"
        self.page_images.append((name, num))
        self.ops.append(b"q %.2f 0 0 %.2f %.2f %.2f cm /%s Do Q" % (dw, dh, x + (w - dw) / 2, y + (h - dh) / 2,
                                                                     name.encode()))
        return True

    def end_page(self):
        content = "\n".encode().join(self.ops)
        c = self.obj(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        xo = b" ".join(b"/%s %d 0 R" % (n.encode(), o) for n, o in self.page_images)
        self.pages.append((c, xo))
        self.ops, self.page_images = [], []

    def bytes(self) -> bytes:
        f1 = self.obj(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
        f2 = self.obj(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
        pages_no = len(self.objs) + len(self.pages) + 1
        kids = []
        for c, xo in self.pages:
            kids.append(self.obj(b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %d %d] /Contents %d 0 R "
                                 b"/Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> /XObject << %s >> >> >>"
                                 % (pages_no, self.W, self.H, c, f1, f2, xo)))
        assert self.obj(b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
            b" ".join(b"%d 0 R" % k for k in kids), len(kids))) == pages_no
        cat = self.obj(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_no)
        out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offs = []
        for i, body in enumerate(self.objs, 1):
            offs.append(len(out))
            out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
        x = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(self.objs) + 1)
        for o in offs:
            out += b"%010d 00000 n \n" % o
        out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(self.objs) + 1, cat, x)
        return bytes(out)


PDF_COLS = [("No.", 28), ("Floor / plan", 150), ("Element", 105), ("Size read", 60), ("What is unclear", 170),
            ("Our comment", 125), ("Consultant reply", 132)]


def _cells(i, n) -> list[str]:
    reply = (n.get("reply") or {}).get("text", "") if isinstance(n.get("reply"), dict) else ""
    return [str(i), _floor(n), f"{n.get('mark') or '(no mark)'} {_what(n)}", n.get("size") or "-",
            n.get("issue") or "-", n.get("text") or "-", reply]


def _row_lines(i, n) -> list[list[str]]:
    return [_wrap(v, w - 8, 8.5)[:8] for v, (_, w) in zip(_cells(i, n), PDF_COLS)]


def write_pdf(project: str, items: list[dict], out: str, sender: str = "") -> str:
    """The same queries as a PDF: two to a page for those with a screenshot - details on
    the left, the screenshot on the right, a box for the consultant's answer - then the
    ones without a screenshot as a table, with a reply column."""
    pdf = _Pdf()
    job = os.path.basename(os.path.normpath(project))
    M = 36.0
    blocks, plain = [], []
    for i, n in enumerate(items, 1):
        imgs = [p for p in (image_path(project, x) for x in n.get("images") or []) if p]
        if not imgs:
            plain.append((i, n))
            continue
        blocks.append((i, n, imgs[0], 0, len(imgs)))
        for k, p in enumerate(imgs[1:], 2):
            blocks.append((i, n, p, k, len(imgs)))
    per_page, top = 2, pdf.H - M - 50
    bh = (top - M) / per_page
    # the table's rows, to a page by their height
    table_pages, page, y = [], [], top - 22
    for i, n in plain:
        h = 12 * max(1, max(len(x) for x in _row_lines(i, n))) + 8
        if page and y - h < M:
            table_pages.append(page)
            page, y = [], top - 22
        page.append((i, n, h))
        y -= h
    if page:
        table_pages.append(page)
    pic_pages = -(-len(blocks) // per_page)
    pages = max(1, pic_pages + len(table_pages))

    def heading(pg):
        pdf.text(M, pdf.H - M - 6, f"{job} - queries for the structural consultant", 15, bold=True)
        pdf.text(M, pdf.H - M - 24, f"Sent {datetime.date.today().strftime('%d %b %Y')}" +
                 (f" by {sender}" if sender else "") + f"  |  {len(items)} quer{'y' if len(items) == 1 else 'ies'}"
                 f"  |  page {pg + 1} of {pages}", 9, grey=True)
        if pg == 0:
            pdf.text(M, pdf.H - M - 38, "Please answer each query (sizes as 230 x 450, or \"see revised drawing\"), "
                                        "or reply in the Excel version of this list.", 9, grey=True)
    if not items:
        heading(0)
        pdf.text(M, top - 30, "No queries.", 11)
        pdf.end_page()
    for pg in range(pic_pages):
        heading(pg)
        for j, (i, n, img, k, nimg) in enumerate(blocks[pg * per_page:(pg + 1) * per_page]):
            y1 = top - j * bh
            y0 = y1 - bh + 10
            pdf.rect(M, y0, pdf.W - 2 * M, y1 - y0, stroke=(0.82, 0.8, 0.75))
            tx, tw = M + 10, 300.0
            y = y1 - 20
            head = f"{i}.  {(n.get('mark') or '(no mark)')}  -  {_what(n)}" + (f"  (screenshot {k} of {nimg})" if k else "")
            pdf.text(tx, y, head, 12, bold=True)
            y -= 16
            if not k:
                for label, val in (("Floor / plan", _floor(n)), ("Size read", n.get("size") or "-"),
                                   ("What is unclear", n.get("issue") or "-"), ("Our comment", n.get("text") or "-"),
                                   ("Raised by", " ".join(x for x in (n.get("by") or "", n.get("date") or "") if x) or "-")):
                    pdf.text(tx, y, label, 8, grey=True)
                    y -= 11
                    for line in _wrap(val, tw, 9.5)[:6]:
                        pdf.text(tx, y, line, 9.5)
                        y -= 12
                    y -= 3
                reply = (n.get("reply") or {}).get("text", "") if isinstance(n.get("reply"), dict) else ""
                bx_h = max(30.0, y - y0 - 14)
                pdf.text(tx, y0 + 8 + bx_h + 2, "Consultant reply", 8, grey=True)
                pdf.rect(tx, y0 + 8, tw, bx_h, fill=(1.0, 0.965, 0.784), stroke=(0.85, 0.75, 0.3))
                ry = y0 + 8 + bx_h - 12
                for line in _wrap(reply, tw - 8, 9)[:int(bx_h // 11)]:
                    pdf.text(tx + 4, ry, line, 9)
                    ry -= 11
            ix, iw = tx + tw + 14, pdf.W - 2 * M - tw - 34
            jpeg = _as_jpeg(img)
            if not (jpeg and pdf.image(jpeg, ix, y0 + 6, iw, y1 - y0 - 12)):
                pdf.text(ix + 10, (y0 + y1) / 2, "(screenshot could not be added)", 9, grey=True)
        pdf.end_page()
    for t, rows_ in enumerate(table_pages):
        heading(pic_pages + t)
        y = top
        x = M
        for name, w in PDF_COLS:                                  # the header row
            pdf.rect(x, y - 18, w, 18, fill=(0.122, 0.435, 0.698), stroke=(0.8, 0.8, 0.8))
            pdf.ops.append(b"1 1 1 rg BT /F2 8.5 Tf %.1f %.1f Td " % (x + 4, y - 12) + _pdf_str(name) + b" Tj ET")
            x += w
        y -= 18
        for i, n, h in rows_:
            x = M
            for (name, w), lines in zip(PDF_COLS, _row_lines(i, n)):
                fill = (1.0, 0.965, 0.784) if name == "Consultant reply" else None
                pdf.rect(x, y - h, w, h, fill=fill, stroke=(0.8, 0.8, 0.8))
                ly = y - 12
                for line in lines:
                    pdf.text(x + 4, ly, line, 8.5)
                    ly -= 12
                x += w
            y -= h
        pdf.end_page()
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "wb") as f:
        f.write(pdf.bytes())
    return out


# ---------------------------------------------------------------- one element in two analyses
def _centre(e):
    g = e.geometry or {}
    if g.get("start") and g.get("end"):
        if g.get("mid"):
            return (float(g["mid"][0]), float(g["mid"][1]))
        return ((g["start"][0] + g["end"][0]) / 2.0, (g["start"][1] + g["end"][1]) / 2.0)
    if g.get("outline"):
        pts = g["outline"]
        return (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts))
    if g.get("x") is not None:
        return (float(g["x"]), float(g["y"]))
    return None


def _seg(e):
    g = e.geometry or {}
    if g.get("start") and g.get("end"):
        return (tuple(g["start"][:2]), tuple(g["end"][:2]))
    return None


def _angle(seg) -> float:
    (x1, y1), (x2, y2) = seg
    return math.degrees(math.atan2(y2 - y1, x2 - x1)) % 180.0


def _to_seg(p, seg) -> float:
    (x1, y1), (x2, y2) = seg
    dx, dy = x2 - x1, y2 - y1
    L2 = dx * dx + dy * dy
    t = 0.0 if L2 == 0 else max(0.0, min(1.0, ((p[0] - x1) * dx + (p[1] - y1) * dy) / L2))
    return math.hypot(p[0] - (x1 + t * dx), p[1] - (y1 + t * dy))


def _read_mark(e) -> str:
    v = e.props.get("mark")
    return str((v.value if v is not None else "") or e.mark or "").strip().upper()


def _kept(e) -> bool:
    return e.type in COMPARED and not e.extra.get("split_into") and not e.extra.get("merged_into") \
        and _centre(e) is not None


def _shift(old: list, new: list):
    """How far the revised plan was moved: the median offset of the marked columns both
    share (a mark used once on each). (0, 0) when they do not agree on one."""
    def once(els):
        c = {}
        for e in els:
            if e.type == "column" and _read_mark(e):
                c.setdefault(_read_mark(e), []).append(e)
        return {k: v[0] for k, v in c.items() if len(v) == 1}
    a, b = once(old), once(new)
    offs = []
    for k in set(a) & set(b):
        pa, pb = _centre(a[k]), _centre(b[k])
        offs.append((pb[0] - pa[0], pb[1] - pa[1]))
    if len(offs) < 3:
        return (0.0, 0.0)
    xs, ys = sorted(o[0] for o in offs), sorted(o[1] for o in offs)
    mx, my = xs[len(xs) // 2], ys[len(ys) // 2]
    near = sum(1 for o in offs if math.hypot(o[0] - mx, o[1] - my) < 50.0)
    return (mx, my) if near >= 0.6 * len(offs) else (0.0, 0.0)


TOL = {"column": 300.0, "beam": 300.0, "beam_span": 300.0, "wall": 300.0, "footing": 400.0, "slab": 1200.0,
       "raft": 1500.0}


def _distance(o, n, shift) -> float | None:
    """How far element n (new) is from o (old, moved by the plan's shift); None when it
    cannot be the same element."""
    if o.type != n.type:
        return None
    co = _centre(o)
    co = (co[0] + shift[0], co[1] + shift[1])
    go, gn = o.geometry or {}, n.geometry or {}
    if go.get("mid") or gn.get("mid"):
        # a curved beam: the same curve when its middle and both ends are where they were
        if not (go.get("mid") and gn.get("mid") and go.get("start") and gn.get("start")):
            return None
        mv = lambda p: (p[0] + shift[0], p[1] + shift[1])      # noqa: E731
        a, b = mv(go["start"]), mv(go["end"])
        c, d_ = gn["start"], gn["end"]
        ends = min(max(math.dist(a, c[:2]), math.dist(b, d_[:2])), max(math.dist(a, d_[:2]), math.dist(b, c[:2])))
        d = max(math.dist(co, _centre(n)), ends)
        return d if d <= TOL.get(o.type, 300.0) else None
    so, sn = _seg(o), _seg(n)
    if so and sn:
        so = ((so[0][0] + shift[0], so[0][1] + shift[1]), (so[1][0] + shift[0], so[1][1] + shift[1]))
        da = abs(_angle(so) - _angle(sn))
        if min(da, 180 - da) > 8.0:
            return None
        # lengthened or shortened in the revision: the two lines must still overlap, side by side
        d = min(_to_seg(co, sn), _to_seg(_centre(n), so))
        if d > TOL.get(o.type, 300.0):
            return None
        return d
    cn = _centre(n)
    d = math.hypot(co[0] - cn[0], co[1] - cn[1])
    return d if d <= TOL.get(o.type, 300.0) else None


def match(old: list, new: list) -> dict:
    """old id -> new id, one to one, for the elements of one plan in two analyses. A pair
    is taken only when it is clearly the nearest for both; a shared mark wins a tie."""
    old = [e for e in old if _kept(e)]
    new = [e for e in new if _kept(e)]
    shift = _shift(old, new)
    pairs = []
    for o in old:
        for n in new:
            d = _distance(o, n, shift)
            if d is None:
                continue
            same = _read_mark(o) and _read_mark(o) == _read_mark(n)
            pairs.append((d - (150.0 if same else 0.0), o.id, n.id))
    pairs.sort()
    used_o, used_n, out = set(), set(), {}
    for score, oid, nid in pairs:
        if oid in used_o or nid in used_n:
            continue
        out[oid] = nid
        used_o.add(oid)
        used_n.add(nid)
    return out


# ---------------------------------------------------------------- plans of two analyses
def _plans(b) -> dict:
    """plan sheet id -> (source drawing path, plan key) for the plans the review shows."""
    out = {}
    for sid, s in b.sheets.items():
        if s.reader != "plan.legacy":
            continue
        view = s.extra.get("view_sheet", "")
        v = b.sheets.get(view)
        src = (v.file if v else "") or s.file
        layout = sid.split(":", 1)[1] if ":" in sid else sid
        title = re.sub(r"\s+", " ", ((v.title if v else "") or s.title or "")).strip().upper()
        out[sid] = (os.path.normcase(os.path.abspath(src)) if src else "", layout, title,
                    s.extra.get("level_key") or s.level or "")
    return out


def pair_plans(old_b, new_b, revised: dict) -> list[tuple[str, str]]:
    """(old plan id, new plan id) for each plan of a revised drawing: revised maps the old
    drawing's path to the new one's (the same path for a revision saved over the old file).
    The plans are paired by their place in the drawing (layout), else their title, else
    their floor - each only when it is the one candidate."""
    po, pn = _plans(old_b), _plans(new_b)
    rev = {os.path.normcase(os.path.abspath(k)): os.path.normcase(os.path.abspath(v)) for k, v in revised.items()}
    out = []
    for oid, (src, layout, title, lv) in po.items():
        to = rev.get(src)
        if not to:
            continue
        cands = [nid for nid, x in pn.items() if x[0] == to]
        for key in (1, 2, 3):
            want = (src, layout, title, lv)[key]
            if not want:
                continue
            hit = [nid for nid in cands if pn[nid][key] == want]
            if len(hit) == 1:
                out.append((oid, hit[0]))
                break
    return out


def _vals(e) -> dict:
    return {p: e.final(p) for p in ("mark",) + SIZE_PROPS if e.final(p) not in (None, "")}


def _same(a, b) -> bool:
    try:
        return abs(float(a) - float(b)) < 0.5
    except (TypeError, ValueError):
        return str(a or "").strip().upper() == str(b or "").strip().upper()


def _diff(vo: dict, vn: dict) -> list[str]:
    return [f"{p} {_n(vo.get(p, '-'))} -> {_n(vn.get(p, '-'))}" for p in sorted(set(vo) | set(vn))
            if not _same(vo.get(p), vn.get(p))]


def as_drawn(path: str, overrides: dict | None):
    """An analysis as the drawings alone give it: linked again without the review's
    decisions (what a mark prefix means is kept), so a revision is compared drawing to
    drawing, never with a decision."""
    from building import Building
    from building.linker import link
    b = Building.load(path)
    link(b, {"prefixes": (overrides or {}).get("prefixes") or {}})
    return b


def _desc(e) -> str:
    return " ".join(x for x in ((e.final("mark") or "(no mark)"), e.type.replace("_", " "), _size_text(e)) if x)


def changes(old_b, new_b, pairs: list) -> list[dict]:
    """What the revision changed, plan by plan: added, removed, changed (mark / size) and
    moved (over 50 mm) elements, each with where to look."""
    out = []
    for oid, nid in pairs:
        old = [e for e in old_b.elements if e.sheet == oid and _kept(e)]
        new = [e for e in new_b.elements if e.sheet == nid and _kept(e)]
        m = match(old, new)
        shift = _shift(old, new)
        by_new = {e.id: e for e in new}
        s = new_b.sheets[nid]
        plan = _plan_title(new_b, s)
        lv = s.extra.get("level_key") or s.level or ""
        for o in old:
            n = by_new.get(m.get(o.id, ""))
            co = _centre(o)
            if n is None:
                out.append({"kind": "removed", "plan": plan, "lv": lv, "sheet": nid, "type": o.type,
                            "mark": o.final("mark") or "", "el": "",
                            "at": [round(co[0] + shift[0]), round(co[1] + shift[1])],
                            "text": f"{_desc(o)} is not in the revised drawing"})
                continue
            diff = _diff(_vals(o), _vals(n))
            cn = _centre(n)
            moved = math.hypot(co[0] + shift[0] - cn[0], co[1] + shift[1] - cn[1])
            so, sn = _seg(o), _seg(n)
            if so and sn:
                lo, ln = math.dist(*so), math.dist(*sn)
                if abs(lo - ln) > 50.0:
                    diff.append(f"length {_n(lo)} -> {_n(ln)}")
            if diff:
                out.append({"kind": "changed", "plan": plan, "lv": lv, "sheet": nid, "type": n.type,
                            "mark": n.final("mark") or "", "el": n.id, "at": [round(cn[0]), round(cn[1])],
                            "text": f"{_desc(o)}: " + ", ".join(diff)})
            elif moved > 50.0:
                out.append({"kind": "moved", "plan": plan, "lv": lv, "sheet": nid, "type": n.type,
                            "mark": n.final("mark") or "", "el": n.id, "at": [round(cn[0]), round(cn[1])],
                            "text": f"{_desc(n)} moved {_n(moved)} mm"})
        got = set(m.values())
        for n in new:
            if n.id not in got:
                cn = _centre(n)
                out.append({"kind": "added", "plan": plan, "lv": lv, "sheet": nid, "type": n.type,
                            "mark": n.final("mark") or "", "el": n.id, "at": [round(cn[0]), round(cn[1])],
                            "text": f"{_desc(n)} is new in the revised drawing"})
    order = {"changed": 0, "added": 1, "removed": 2, "moved": 3}
    out.sort(key=lambda c: (c["lv"], c["plan"], order[c["kind"]], c["mark"]))
    return out


# ---------------------------------------------------------------- after an analysis
def _runs_before(project: str, run_dir: str) -> list[str]:
    runs = os.path.join(project, "runs")
    me = os.path.basename(os.path.normpath(run_dir))
    try:
        names = sorted(d for d in os.listdir(runs) if d < me)
    except OSError:
        return []
    return [os.path.join(runs, d) for d in reversed(names)
            if os.path.exists(os.path.join(runs, d, "building.json"))
            and os.path.exists(os.path.join(runs, d, "review.html"))]


def _seen(project: str) -> dict:
    try:
        with open(folder(project, "drawings_seen.json"), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def _stamp(path: str):
    try:
        st = os.stat(path)
        return [round(st.st_mtime, 1), st.st_size]
    except OSError:
        return None


def revised_drawings(project: str, files: list[str], revisions: list[dict]) -> dict:
    """old drawing path -> new path, for this analysis: a "Replace with revision" since the
    last analysis, or a drawing saved over with a new version (same path, new date)."""
    seen = _seen(project)
    out = {}
    for r in revisions or []:
        if r.get("new") in files and not r.get("analysed"):
            out[r["old"]] = r["new"]
    for f in files:
        was = seen.get(os.path.normcase(os.path.abspath(f)))
        now = _stamp(f)
        if was and now and was != now and f not in out.values():
            out[f] = f
    return out


def after_analysis(project: str, run_dir: str, files: list[str], revisions: list[dict] | None = None,
                   log=print) -> dict:
    """Decisions and notes follow their elements into this analysis, and a revised
    drawing's changes are listed. Returns {"moved": n, "relink": bool, ...}."""
    from building import Building
    res = {"moved": 0, "relink": False, "notes_moved": 0, "unmatched": [], "changes": 0}
    revised = revised_drawings(project, files, revisions or [])
    new_path = os.path.join(run_dir, "building.json")
    if not os.path.exists(new_path):
        return res
    ov_path = os.path.join(project, "overrides.json")
    try:
        with open(ov_path, encoding="utf-8") as f:
            ov = json.load(f)
    except (OSError, ValueError):
        ov = {}
    nd = load(project)
    wanted = set((ov.get("elements") or {}).keys()) | set(ov.get("exclude") or []) | \
        {n["el"] for n in nd["notes"] if n.get("el")}
    new_b = Building.load(new_path)
    new_by = {e.id: e for e in new_b.elements}
    before = _runs_before(project, run_dir)
    old_b = None
    drawn_by = ({}, {})                 # the two analyses as the drawings alone give them

    def old():
        nonlocal old_b
        if old_b is None and before:
            old_b = Building.load(os.path.join(before[0], "building.json"))
        return old_b

    # which of them are no longer the same element in this analysis
    lost = []
    for eid in sorted(wanted):
        n = new_by.get(eid)
        if n is None:
            lost.append(eid)
            continue
        if revised and old() is not None:
            o = next((e for e in old_b.elements if e.id == eid), None)
            if o is not None and _kept(o) and _kept(n) and _distance(o, n, (0.0, 0.0)) is None and \
                    _distance(o, n, _shift([x for x in old_b.elements if x.sheet == o.sheet],
                                           [x for x in new_b.elements if x.sheet == n.sheet])) is None:
                lost.append(eid)        # the id is reused by a different element of the revision
    moved = {}
    if (lost or revised) and old() is not None:
        old_by = {e.id: e for e in old_b.elements}
        # plans paired: a revised drawing's by its pairing, every other plan with itself
        pairs = pair_plans(old_b, new_b, revised) if revised else []
        paired = dict(pairs)
        for sid in {old_by[x].sheet for x in lost if x in old_by}:
            if sid not in paired and sid in new_b.sheets:
                paired[sid] = sid
        maps = {}
        for eid in lost:
            o = old_by.get(eid)
            if o is None or o.sheet not in paired:
                continue
            key = (o.sheet, paired[o.sheet])
            if key not in maps:
                maps[key] = match([e for e in old_b.elements if e.sheet == key[0]],
                                  [e for e in new_b.elements if e.sheet == key[1]])
            to = maps[key].get(eid)
            if to and to != eid:
                moved[eid] = to
        if revised and pairs:
            try:
                old_d = as_drawn(os.path.join(before[0], "building.json"), ov)
                new_d = as_drawn(new_path, ov)
                drawn_by = ({e.id: e for e in old_d.elements}, {e.id: e for e in new_d.elements})
            except Exception as exc:          # compared as analysed, decisions and all
                log(f"(the revision is compared with the decisions in: {exc})")
                old_d, new_d = old_b, new_b
            ch = changes(old_d, new_d, pairs)
            res["changes"] = len(ch)
            with open(os.path.join(run_dir, "revision_changes.json"), "w", encoding="utf-8") as f:
                json.dump({"from": os.path.basename(before[0]), "drawings": revised, "changes": ch}, f, indent=1)
            log(f"Revised drawing(s): {len(ch)} change(s) found - listed in Review (Changes in the revision).")
    res["unmatched"] = [x for x in lost if x not in moved]
    # decisions moved to their element's new id (the old id is not kept: in a revision it
    # can belong to another element now)
    if moved:
        els = dict(ov.get("elements") or {})
        for a, b in moved.items():
            if a in els:
                props = els.pop(a)
                els[b] = {**props, **els.get(b, {})}
                res["moved"] += 1
        exc = list(ov.get("exclude") or [])
        for i, x in enumerate(exc):
            if x in moved:
                exc[i] = moved[x]
                res["moved"] += 1
        if res["moved"]:
            ov["elements"] = els
            ov["exclude"] = list(dict.fromkeys(exc))
            with open(ov_path, "w", encoding="utf-8") as f:
                json.dump(ov, f, indent=1)
            res["relink"] = True
            log(f"{res['moved']} decision(s) moved to their element in the new analysis.")
            _log_carried(project, run_dir, moved)
    # notes: follow their element; in a revised drawing, what became of what they asked
    rev_src = {os.path.normcase(os.path.abspath(k)) for k in revised}
    old_plans = _plans(old_b) if (revised and old_b is not None) else {}
    old_by = {e.id: e for e in old_b.elements} if old_b is not None else {}
    with _lock:
        nd = load(project)
        for n in nd["notes"]:
            eid = n.get("el")
            if not eid:
                continue
            o = old_by.get(eid)
            to = moved.get(eid, eid if eid in new_by and eid not in lost else None)
            if to and to != eid:
                n["el"] = to
                n["sheet"] = new_by[to].sheet
                res["notes_moved"] += 1
            # a note on a drawing that was revised: what became of the element it asks about
            if o is None or old_plans.get(o.sheet, ("",))[0] not in rev_src:
                continue
            if to is None:
                n["revision"] = {"result": "not found", "date": _now(),
                                 "text": "Not found in the revised drawing - check the plan."}
            else:
                diff = _diff(_vals(drawn_by[0].get(eid, o)), _vals(drawn_by[1].get(to, new_by[to])))
                n["revision"] = {"result": "changed" if diff else "same", "date": _now(),
                                 "text": ("Changed in the revised drawing: " + ", ".join(diff)) if diff
                                 else "The revised drawing shows it as before."}
        save(project, nd)
    if res["notes_moved"]:
        log(f"{res['notes_moved']} note(s) moved to their element in the new analysis.")
    if res["unmatched"]:
        log(f"{len(res['unmatched'])} decision(s) or note(s) could not be matched to an element of the new "
            f"analysis - kept as they were.")
    remember_drawings(project, files)
    return res


def remember_drawings(project: str, files: list[str]) -> None:
    seen = _seen(project)
    for f in files:
        st = _stamp(f)
        if st:
            seen[os.path.normcase(os.path.abspath(f))] = st
    os.makedirs(folder(project), exist_ok=True)
    with open(folder(project, "drawings_seen.json"), "w", encoding="utf-8") as f:
        json.dump(seen, f, indent=1)


def _log_carried(project: str, run_dir: str, moved: dict) -> None:
    p = folder(project, "carried.json")
    try:
        with open(p, encoding="utf-8") as f:
            d = json.load(f)
    except (OSError, ValueError):
        d = []
    d.append({"run": os.path.basename(os.path.normpath(run_dir)), "date": _now(), "moved": moved})
    os.makedirs(folder(project), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(d, f, indent=1)


def revision_changes(run_dir: str) -> dict | None:
    try:
        with open(os.path.join(run_dir or "", "revision_changes.json"), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("project")
    ap.add_argument("what", choices=["report"])
    ap.add_argument("out")
    a = ap.parse_args(argv)
    items = rows(a.project)
    (write_pdf if a.out.lower().endswith(".pdf") else write_xlsx)(a.project, items, a.out)
    print(f"Wrote {a.out} ({len(items)} queries)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
