r"""Build the structural schedule workbook the consultant fills in.

    python make_schedules.py                       blank template
    python make_schedules.py --from ..\extracted   pre-filled with everything the
                                                   analysis read from the drawing
    python make_schedules.py --from F1 F2 F3       one workbook for several floors,
                                                   one analysis folder per floor
    python make_schedules.py --out MyFile.xlsx

One workbook, one sheet per element type, with the columns Revit needs. There are
deliberately no formulas: every cell is either a heading or something a person
types, so nothing can be silently broken by an edit.

Pre-filled values are what the analysis read off the drawing: marks, sizes, beam
levels, slab thickness and sunk depth, sunk zones, openings and slab folds. They
are shaded grey so the consultant can see what to check; whatever is still yellow
the drawing did not give. Where one mark appears in the drawing with more than one
value, the commonest is used and the Remark lists the others.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
INK = "1F2A37"
BLUE = "1F6FB2"
HEAD_FILL = PatternFill("solid", fgColor="16202B")
FILL_ME = PatternFill("solid", fgColor="FFF6D6")      # yellow: you fill this in
REF_FILL = PatternFill("solid", fgColor="EDEBE5")     # grey: reference, do not edit
EG_FILL = PatternFill("solid", fgColor="F4F3EF")
THIN = Side(style="thin", color="D5D2CA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GRADES = "M20,M25,M30,M35,M40,M45,M50"
BEAM_TYPES = "B,BC,BP,EB,HB,MLB,GB,BKT"
BEAM_TYPE = re.compile(r"(BKT|MLB|GB|BC|BP|EB|HB|B)")
BEAM_PREFIX = re.compile(BEAM_TYPE.pattern + r"(\d+[A-Z]?)")
LEVEL_REF = "=Levels!$A$6:$A$45"   # Levels data rows: 6 to 45

# ---------------------------------------------------------------- sheet specs
# (header, width, required, validation-list or None)
SHEETS: dict[str, dict] = {
    "Levels": {
        "note": "Fill this in first. Every other sheet refers to these level names.",
        "cols": [
            ("Level Name", 22, True, None),
            ("Structural Level SSL (mm)", 26, True, None),
            ("Floor to Floor (mm)", 20, False, None),
            ("Datum Note", 40, False, None),
        ],
        "example": ["1st Floor", 3300, 3000, "Levels are to project datum, 0 = plinth top"],
        "seed": [["Plinth", None, None, None], ["1st Floor", None, None, None],
                 ["2nd Floor", None, None, None], ["Terrace", None, None, None]],
    },
    "Columns": {
        "note": "One row per column mark. Split grouped rows such as 10,14,22,26 into one row each.",
        "cols": [
            ("Mark", 12, True, None),
            ("CAD tag", 12, False, None),
            ("Type", 10, True, "C,FC"),
            ("Width (mm)", 12, True, None),
            ("Depth (mm)", 12, True, None),
            ("Base Level", 16, True, LEVEL_REF),
            ("Top Level", 16, True, LEVEL_REF),
            ("Concrete Grade", 15, True, GRADES),
            ("Main Steel", 26, False, None),
            ("Links", 22, False, None),
            ("Continues Above", 16, False, "Y,N"),
            ("Remark", 34, False, None),
        ],
        "example": ["C-3", "3", "C", 230, 1200, "Plinth", "1st Floor", "M30",
                    "12-T16", "T8 @ 150 c/c", "Y", ""],
    },
    "Beams": {
        "note": "Beams and girders together. Give every depth: a blank depth cannot be modelled.",
        "cols": [
            ("Mark", 12, True, None),
            ("CAD tag", 14, False, None),
            ("Type", 10, True, BEAM_TYPES),
            ("Width (mm)", 12, True, None),
            ("Depth (mm)", 12, True, None),
            ("Level", 16, True, LEVEL_REF),
            ("Top of Beam Offset (mm)", 22, False, None),
            ("Concrete Grade", 15, True, GRADES),
            ("Top Steel", 24, False, None),
            ("Bottom Steel", 24, False, None),
            ("Extra Top", 22, False, None),
            ("Stirrups", 22, False, None),
            ("Sunk Depth (mm)", 16, False, None),
            ("Remark", 34, False, None),
        ],
        "example": ["B-34", "B34", "B", 230, 700, "1st Floor", -100, "M30",
                    "3-T16", "3-T16", "2-T12", "T8 @ 150 c/c", 100, "Partly sunk, see section"],
    },
    "Slabs": {
        "note": "One row per slab mark. A mark used on several panels still gets one row here.",
        "cols": [
            ("Mark", 12, True, None),
            ("CAD tag", 12, False, None),
            ("Thickness (mm)", 15, True, None),
            ("Level", 16, True, LEVEL_REF),
            ("Top Offset (mm)", 16, False, None),
            ("Concrete Grade", 15, True, GRADES),
            ("Span Type", 14, False, "One Way,Two Way"),
            ("Sunk Depth (mm)", 16, False, None),
            ("Top Steel", 24, False, None),
            ("Bottom Steel", 24, False, None),
            ("Remark", 34, False, None),
        ],
        "example": ["S-4A", "S4A", 150, "1st Floor", -150, "M30", "Two Way", 150,
                    "T8 @ 150 c/c", "T8 @ 150 c/c", "Toilet sunk"],
    },
    "Walls": {
        "note": "Lift walls, RCC pardi and retaining walls.",
        "cols": [
            ("Mark", 12, True, None),
            ("CAD tag", 12, False, None),
            ("Thickness (mm)", 15, True, None),
            ("Base Level", 16, True, LEVEL_REF),
            ("Top Level", 16, True, LEVEL_REF),
            ("Concrete Grade", 15, True, GRADES),
            ("Vertical Steel", 26, False, None),
            ("Horizontal Steel", 26, False, None),
            ("Remark", 34, False, None),
        ],
        "example": ["LW-1", "LW1", 230, "Plinth", "Terrace", "M30",
                    "T12 @ 150 c/c", "T8 @ 100 c/c", "Both faces"],
    },
    "Sunk and Openings": {
        "note": "Every sunk zone, opening and cut-out, keyed to the slab panel it sits in.",
        "cols": [
            ("ID", 10, True, None),
            ("Type", 14, True, "Sunk,Opening,Cut-out,Pool,Deck"),
            ("Depth (mm)", 13, True, None),
            ("Level", 16, True, LEVEL_REF),
            ("In Slab Panel", 16, False, None),
            ("Where", 30, False, None),
            ("Remark", 34, False, None),
        ],
        "example": ["SK-1", "Sunk", 150, "1st Floor", "S-4A", "Toilet, grid C-D / 3-4", ""],
    },
    "Slab Folds": {
        "note": "The downstand where a slab steps down to a sunk slab (the paired red lines "
                "on the slab layer). Revit builds these from the drawing; this sheet is for checking.",
        "cols": [
            ("ID", 10, True, None),
            ("Width (mm)", 12, True, None),
            ("Depth (mm)", 12, True, None),
            ("Level", 16, True, LEVEL_REF),
            ("Top Offset (mm)", 16, False, None),
            ("Where", 30, False, None),
            ("Remark", 40, False, None),
        ],
        "example": ["F-1", 200, 250, "1st Floor", 0, "x 12400, y 5300 from base point",
                    "between S7 and S4, step 100 mm"],
    },
    "Stairs": {
        "note": "Stair flights and landings.",
        "cols": [
            ("Mark", 12, True, None),
            ("CAD tag", 12, False, None),
            ("Waist (mm)", 13, True, None),
            ("Level From", 16, True, LEVEL_REF),
            ("Level To", 16, True, LEVEL_REF),
            ("Concrete Grade", 15, True, GRADES),
            ("Main Steel", 24, False, None),
            ("Remark", 34, False, None),
        ],
        "example": ["ST-1", "ST1", 150, "Plinth", "1st Floor", "M30", "T12 @ 150 c/c", ""],
    },
    "Revisions": {
        "note": "One row per issue. This is what tells us which floors need re-running.",
        "cols": [
            ("Rev", 8, True, None),
            ("Date", 14, True, None),
            ("Issued By", 18, False, None),
            ("What Changed", 52, True, None),
            ("Marks Affected", 30, False, None),
        ],
        "example": ["R1", "2026-07-14", "Struct. consultant",
                    "Beam depths revised on the north wing", "B-54, B-55, B-58"],
    },
}

SHEET_ORDER = ["Read Me", "Levels", "Columns", "Beams", "Slabs", "Walls",
               "Sunk and Openings", "Slab Folds", "Stairs", "Revisions"]


# ---------------------------------------------------------------- mark tidying
def normalise(tag: str, kind: str) -> str:
    """Turn a CAD tag into a schedule mark: '12' -> 'C-12', 'B107' -> 'B-107'."""
    t = (tag or "").strip().upper()
    if not t:
        return ""
    t = t.split("/")[0]                       # compound tags keep their first part
    if kind == "column":
        if t.startswith("FC"):
            return "FC-" + t[2:]
        if re.fullmatch(r"\d+[A-Z]?", t):
            return "C-" + t
        return t
    if kind == "wall":
        return "LW-" + t[2:] if t.startswith("LW") else t
    if kind == "beam":
        m = BEAM_PREFIX.fullmatch(t)
        return f"{m.group(1)}-{m.group(2)}" if m else t
    if kind == "slab":
        m = re.fullmatch(r"(ST|S)(\d+[A-Z]?)", t)
        return f"{m.group(1)}-{m.group(2)}" if m else t
    return t


def _sortkey(mark: str):
    m = re.match(r"([A-Z]+)-?(\d+)([A-Z]?)", mark or "")
    return (m.group(1), int(m.group(2)), m.group(3)) if m else (mark or "", 0, "")


def r5(v):
    """Drawing sizes to the nearest 5 mm: 1499.9 is a 1500 column."""
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return int(round(v / 5.0) * 5) if v else None


def commonest(values, label="", unit=" mm"):
    """The commonest value, and a remark naming the others if there are any.

    Values of None are ignored. A mark drawn at two sizes is a question for the
    consultant, not something to average away, so the remark says so.
    """
    c = Counter(v for v in values if v is not None)
    if not c:
        return None, ""
    best = c.most_common(1)[0][0]
    if len(c) == 1:
        return best, ""
    shown = ", ".join(f"{v}{unit} x{n}" for v, n in c.most_common())
    return best, f"Drawing shows {label}{shown}"


def _load(folder: str) -> dict:
    """model.json if the folder has one, else what the CSVs can give.

    model.json is the full record (levels, folds, sunk zones, openings); the CSVs
    are the fallback for an older extraction that left only those behind.
    """
    p = os.path.join(folder, "model.json")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)

    def rows(name):
        q = os.path.join(folder, name)
        if not os.path.exists(q):
            return []
        with open(q, encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))

    def num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    return {
        "columns": [{"mark": r.get("mark"), "kind": r.get("type"),
                     "width": num(r.get("width_mm")), "depth": num(r.get("depth_mm"))}
                    for r in rows("columns.csv")],
        "beams": [{"mark": r.get("mark"), "kind": r.get("type"),
                   "width": num(r.get("width_mm")), "depth": num(r.get("depth_mm"))}
                  for r in rows("beams.csv")],
        "slab_tags": [{"mark": r.get("mark"), "thickness": num(r.get("thickness_mm"))}
                      for r in rows("slab_tags.csv")],
        "walls": [{"mark": r.get("mark"), "thickness": num(r.get("thickness_mm"))}
                  for r in rows("walls.csv")],
    }


def _inside(pt, poly) -> bool:
    x, y = pt
    hit = False
    for i in range(len(poly)):
        x1, y1 = poly[i][0], poly[i][1]
        x2, y2 = poly[i - 1][0], poly[i - 1][1]
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            hit = not hit
    return hit


def _centre(outline):
    pts = outline or []
    if not pts:
        return None
    return (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts))


def drawing_meta(folder: str, data: dict) -> dict:
    """Level, wing and drawing date, for the Levels sheet and the file name."""
    src = os.path.basename(data.get("source_file") or "")
    wing = re.search(r"\(\s*([A-Z])\s+WING\s*\)", src.upper())
    d = re.search(r"(\d{2})-(\d{2})-(\d{4})", src) or re.match(r"(\d{2})_(\d{2})_(\d{4})", src)
    return {
        "folder": folder,
        "source": src,
        "level": (data.get("level_name") or "").strip(),
        "wing": wing.group(1) if wing else "",
        "date": f"{d.group(3)}-{d.group(2)}-{d.group(1)}" if d else "",
    }


def read_analysis(folder: str, ids: Counter) -> tuple[dict[str, list[dict]], dict]:
    """Every row the analysis can fill, sheet by sheet, for one floor's folder.

    A row is {column heading: value}. `ids` numbers sunk zones, openings and folds
    across the whole workbook so every ID is unique.
    """
    data = _load(folder)
    meta = drawing_meta(folder, data)
    lv = meta["level"] or None
    bx, by = (data.get("base_point") or [0.0, 0.0])[:2]

    def where(pt):
        return f"x {pt[0] - bx:.0f}, y {pt[1] - by:.0f} from base point" if pt else None

    groups: dict[str, dict[str, list]] = {s: {} for s in ("Columns", "Beams", "Slabs", "Walls", "Stairs")}

    def add(sheet, mark, item):
        if mark:
            groups[sheet].setdefault(mark, []).append(item)

    for c in data.get("columns", []):
        tag = (c.get("mark") or "").strip()
        up = tag.upper()
        if up.startswith("LW"):
            add("Walls", normalise(tag, "wall"), {"tag": tag, "thickness": r5(min(c.get("width") or 0, c.get("depth") or 0) or None)})
        elif up.startswith("ST"):
            add("Stairs", normalise(tag, "slab"), {"tag": tag})
        else:
            add("Columns", normalise(tag, "column"), dict(c, tag=tag))
    for w in data.get("walls", []):
        tag = (w.get("mark") or "").strip()
        add("Walls", normalise(tag, "wall"), {"tag": tag, "thickness": r5(w.get("thickness"))})
    for b in data.get("beams", []):
        tag = (b.get("mark") or "").strip()
        add("Beams", normalise(tag, "beam"), dict(b, tag=tag))
    for t in data.get("slab_tags", []):
        tag = (t.get("mark") or "").strip()
        add("Stairs" if tag.upper().startswith("ST") else "Slabs", normalise(tag, "slab"), dict(t, tag=tag))

    out: dict[str, list[dict]] = {}

    def tags(items):
        return ", ".join(sorted({i["tag"] for i in items if i.get("tag")}))

    def joined(*remarks):
        return "; ".join(r for r in remarks if r) or None

    rows = []
    for mark, items in groups["Columns"].items():
        # narrow side first, as a schedule writes it; Revit keeps the drawn orientation
        sizes = []
        for i in items:
            a, b = r5(i.get("width")), r5(i.get("depth"))
            sizes.append("%dx%d" % tuple(sorted((a, b))) if a and b else None)
        size, note = commonest(sizes, "sizes ", unit="")
        w, d = (int(v) for v in size.split("x")) if size else (None, None)
        floating = mark.startswith("FC") or any(i.get("kind") == "column_float" for i in items)
        rows.append({"Mark": mark, "CAD tag": tags(items), "Type": "FC" if floating else "C",
                     "Width (mm)": w, "Depth (mm)": d,
                     "Base Level": lv if floating else None, "Top Level": None if floating else lv,
                     "Remark": joined(note, f"{len(items)} on this plan" if len(items) > 1 else "")})
    out["Columns"] = rows

    rows = []
    for mark, items in groups["Beams"].items():
        w, wn = commonest([r5(i.get("width")) for i in items], "widths ")
        d, dn = commonest([r5(i.get("depth")) for i in items], "depths ")
        top, tn = commonest([int(round(i.get("top_mm") or 0)) for i in items], "tops ")
        m = BEAM_TYPE.match(mark)            # 'HB' and 'BKT' are drawn with no number
        compound = sorted({i["tag"] for i in items if "/" in (i.get("tag") or "")})
        rows.append({"Mark": mark, "CAD tag": tags(items), "Type": m.group(1) if m else None,
                     "Width (mm)": w, "Depth (mm)": d, "Level": lv,
                     "Top of Beam Offset (mm)": top if lv or top else None,
                     "Remark": joined(wn, dn, tn, ("tagged " + ", ".join(compound)) if compound else "",
                                      "" if d else "depth not in the drawn schedule")})
    out["Beams"] = rows

    rows = []
    for mark, items in groups["Slabs"].items():
        th, thn = commonest([r5(i.get("thickness")) for i in items], "thicknesses ")
        # an older extraction has no sunk_mm at all: leave it blank, not 0
        sk, skn = commonest([int(round(i.get("sunk_mm") or 0)) if "sunk_mm" in i else None
                             for i in items], "sunk depths ")
        rows.append({"Mark": mark, "CAD tag": tags(items), "Thickness (mm)": th, "Level": lv,
                     "Sunk Depth (mm)": sk, "Remark": joined(thn, skn)})
    out["Slabs"] = rows

    rows = []
    for mark, items in groups["Walls"].items():
        th, thn = commonest([i.get("thickness") for i in items], "thicknesses ")
        rows.append({"Mark": mark, "CAD tag": tags(items), "Thickness (mm)": th,
                     "Top Level": lv, "Remark": joined(thn)})
    out["Walls"] = rows

    out["Stairs"] = [{"Mark": mark, "CAD tag": tags(items), "Level To": lv}
                     for mark, items in groups["Stairs"].items()]

    for sheet in out:
        out[sheet].sort(key=lambda r: _sortkey(r["Mark"]))

    # sunk zones and openings, each placed in the slab panel it sits in
    panels = [(s.get("mark"), s.get("outline")) for s in data.get("slabs", []) if s.get("outline")]

    def panel_of(pt):
        if not pt:
            return None
        for mark, poly in panels:
            if mark and _inside(pt, poly):
                return normalise(mark, "slab")
        return None

    rows = []
    for z in data.get("sunk", []):
        ids["SK"] += 1
        pt = _centre(z.get("outline"))
        src = z.get("depth_source") or ""
        rows.append({"ID": f"SK-{ids['SK']}", "Type": "Sunk", "Depth (mm)": r5(z.get("depth_mm")),
                     "Level": lv, "In Slab Panel": panel_of(pt), "Where": where(pt),
                     "Remark": joined(f"depth from {src}" if src else "",
                                      "" if z.get("depth_mm") else "depth not in the legend")})
    for o in data.get("openings", []):
        ids["OP"] += 1
        pt = _centre(o.get("outline"))
        rows.append({"ID": f"OP-{ids['OP']}", "Type": "Opening", "Level": lv,
                     "In Slab Panel": panel_of(pt), "Where": where(pt)})
    out["Sunk and Openings"] = rows

    rows = []
    for f in data.get("folds", []):
        ids["F"] += 1
        pt = None
        if f.get("start") and f.get("end"):
            pt = ((f["start"][0] + f["end"][0]) / 2.0, (f["start"][1] + f["end"][1]) / 2.0)
        rows.append({"ID": f"F-{ids['F']}", "Width (mm)": r5(f.get("width")),
                     "Depth (mm)": r5(f.get("depth")), "Level": lv,
                     "Top Offset (mm)": int(round(f.get("top_mm") or 0)), "Where": where(pt),
                     "Remark": joined("; ".join(f.get("notes") or []),
                                      "" if f.get("depth") else "no slab on one side, not built")})
    out["Slab Folds"] = rows
    return out, meta


# ---------------------------------------------------------------- the workbook
def style_header(ws, row, cols):
    for i, (head, width, required, _) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=head + (" *" if required else ""))
        c.font = Font(name=FONT, size=10, bold=True, color="F7F6F2")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def build_sheet(wb, name, spec, seed_marks):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False

    t = ws.cell(row=1, column=1, value=name.upper())
    t.font = Font(name=FONT, size=13, bold=True, color=INK)
    n = ws.cell(row=2, column=1, value=spec["note"] + "   Columns marked * must be filled in.")
    n.font = Font(name=FONT, size=9, italic=True, color="5F6C77")

    cols = spec["cols"]
    head_row = 4
    style_header(ws, head_row, cols)

    eg_row = head_row + 1
    for i, v in enumerate(spec["example"], start=1):
        c = ws.cell(row=eg_row, column=i, value=v)
        c.font = Font(name=FONT, size=10, italic=True, color="8A94A0")
        c.fill = EG_FILL
        c.border = BOX
    tagcol = len(cols) + 1
    tag = ws.cell(row=eg_row, column=tagcol, value="<-- EXAMPLE ROW, DELETE BEFORE ISSUE")
    tag.font = Font(name=FONT, size=9, bold=True, color="A4500F")

    seeded = seed_marks.get(name)
    if not seeded:
        heads = [h for h, *_ in cols]
        seeded = [dict(zip(heads, row)) for row in spec.get("seed") or []]
    first_data = eg_row + 1
    n_rows = max(len(seeded) + 15, 40)   # spare rows for marks added later

    for r in range(first_data, first_data + n_rows):
        for i, (head, _w, required, _v) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT, size=10, color=INK)
            c.border = BOX
            if required:
                c.fill = FILL_ME

    # a row is {heading: value}; whatever the drawing gave is shaded grey, so the
    # cells still yellow are exactly what the consultant has left to supply
    from_drawing = name not in ("Levels",)
    for j, item in enumerate(seeded):
        r = first_data + j
        for i, (head, *_rest) in enumerate(cols, start=1):
            v = item.get(head)
            if v in (None, ""):
                continue
            c = ws.cell(row=r, column=i, value=v)
            if from_drawing:
                c.fill = REF_FILL
                c.font = Font(name=FONT, size=10, color="5F6C77")

    last = first_data + n_rows - 1
    for i, (_h, _w, _req, valid) in enumerate(cols, start=1):
        if not valid:
            continue
        formula = valid if valid.startswith("=") else f'"{valid}"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showDropDown=False)
        ws.add_data_validation(dv)
        letter = get_column_letter(i)
        dv.add(f"{letter}{first_data}:{letter}{last}")

    ws.freeze_panes = ws.cell(row=first_data, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(cols))}{last}"
    return ws


def build_readme(wb, project, seeded_from):
    ws = wb.create_sheet("Read Me", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96

    def line(r, a, b="", *, bold=False, size=10, colour=INK, gap=False):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = Font(name=FONT, size=size, bold=bold, color=colour)
        ca.alignment = Alignment(vertical="top")
        cb = ws.cell(row=r, column=2, value=b)
        cb.font = Font(name=FONT, size=size, color=INK)
        cb.alignment = Alignment(vertical="top", wrap_text=True)
        if gap:
            ws.row_dimensions[r].height = 22
        return r + 1

    r = 1
    c = ws.cell(row=r, column=1, value="STRUCTURAL SCHEDULE - REQUIRED FORMAT")
    c.font = Font(name=FONT, size=15, bold=True, color=INK)
    r += 2
    r = line(r, "Project", project)
    r = line(r, "Issued by", "MOFA Studios Pvt. Ltd.")
    r = line(r, "For", "The structural consultant")
    r = line(r, "Version", f"1.0, {date.today():%d %B %Y}")
    for k, src in enumerate(seeded_from):
        r = line(r, "Pre-filled from" if k == 0 else "", src)
    r += 1

    r = line(r, "WHY", "", bold=True)
    r = line(r, "", "The Revit model is generated from your drawing and this workbook. "
                    "The drawn schedule on the sheet stays exactly as it is; this is a "
                    "parallel delivery, not a replacement.")
    r += 1

    r = line(r, "HOW TO FILL IT IN", "", bold=True)
    for item in (
        "Fill the Levels sheet first. Every other sheet refers to those level names.",
        "One row per mark. Split grouped entries such as B-2,4,7 or C-10,14,22,26 into one row each.",
        "One value per cell. No units in the cell, no ranges, no 'as above', no merged cells.",
        "Numbers as numbers, not text. Write 230, not '230 mm' or '230THK'.",
        "Do not add, rename, reorder or delete columns. Add rows freely.",
        "Leave a cell blank only if it genuinely does not apply. A blank Depth cannot be modelled.",
        "Delete the grey example row on each sheet before issuing.",
    ):
        r = line(r, "", "•  " + item)
    r += 1

    r = line(r, "COLOURS", "", bold=True)
    for cell_colour, fill, text in (
        ("Yellow", FILL_ME, "You fill this in. Marked * in the header."),
        ("Grey", REF_FILL, "Read from your CAD drawing by us: marks, sizes, beam levels, slab "
                           "thickness and sunk depths. Check it and overwrite anything wrong."),
        ("White", None, "Optional, but fill it in if you have it."),
    ):
        ca = ws.cell(row=r, column=1, value="   " + cell_colour)
        ca.font = Font(name=FONT, size=10, color=INK)
        if fill:
            ca.fill = fill
        ca.border = BOX
        cb = ws.cell(row=r, column=2, value=text)
        cb.font = Font(name=FONT, size=10, color=INK)
        r += 1
    r += 1

    r = line(r, "THE SHEETS", "", bold=True)
    for name in SHEET_ORDER[1:]:
        r = line(r, "   " + name, SHEETS[name]["note"])
    r += 1

    r = line(r, "FILE NAMING", "", bold=True)
    r = line(r, "", "J1153_<Wing>_<Level>_STR_<YYYY-MM-DD>.xlsx     "
                    "for example  J1153_C_01_STR_2026-07-14.xlsx")
    r += 1
    r = line(r, "NOTE", "", bold=True)
    r = line(r, "", "This workbook contains no formulas. Every cell is either a heading "
                    "or something a person types, so nothing can be broken by an edit.")
    r = line(r, "", "Full requirements: Structural CAD Standard for Automated Revit "
                    "Modelling, issued with this workbook.")
    return ws


def build(out_path: str, seed_marks: dict, project: str, seeded_from: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_ORDER[1:]:
        build_sheet(wb, name, SHEETS[name], seed_marks)
    build_readme(wb, project, seeded_from)
    wb.active = 0
    wb.save(out_path)
    return out_path


def default_name(metas: list[dict]) -> str:
    """J1153_<Wing>_<Level>_STR_<drawing date>.xlsx, as the Read Me asks for."""
    wings = sorted({m["wing"] for m in metas if m["wing"]})
    nums = []
    for m in metas:
        n = re.search(r"\d+", m["level"]) or re.search(r"(\d+)\s*(?:ST|ND|RD|TH)\s+FLOOR", m["source"].upper())
        if n:
            nums.append(int(n.group(1) if n.re.groups else n.group(0)))
    lv = (f"{min(nums):02d}" if len(set(nums)) == 1 else f"{min(nums):02d}-{max(nums):02d}") if nums else "XX"
    dates = sorted(m["date"] for m in metas if m["date"])
    return f"J1153_{''.join(wings) or 'X'}_{lv}_STR_{dates[-1] if dates else date.today():}.xlsx"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--from", dest="src", nargs="+",
                    help="one or more analysis output folders, one per floor")
    ap.add_argument("--out", help="output .xlsx path")
    ap.add_argument("--project", default="J-1153, Reis Magos Goa Housing")
    a = ap.parse_args(argv)

    seed: dict[str, list] = {}
    seeded_from: list[str] = []
    metas: list[dict] = []
    if a.src:
        ids: Counter = Counter()
        for folder in a.src:
            if not os.path.isdir(folder):
                print(f"No such folder: {folder}", file=sys.stderr)
                return 1
            rows, meta = read_analysis(folder, ids)
            metas.append(meta)
            for sheet, rs in rows.items():
                seed.setdefault(sheet, []).extend(rs)
            seeded_from.append(f"{os.path.abspath(folder)}   ({meta['level'] or 'no level'}: "
                               f"{meta['source'] or 'source not recorded'})")
            print(f"  {os.path.abspath(folder)}  level={meta['level'] or '-'}")
            for sheet, rs in rows.items():
                print(f"     {sheet:18s} {len(rs)} rows")
        levels = []
        for m in metas:
            if m["level"] and m["level"] not in levels:
                levels.append(m["level"])
        if levels:
            seed["Levels"] = [{"Level Name": lv} for lv in levels]

    out = a.out or (default_name(metas) if a.src else "Structural-Schedule-Template.xlsx")
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), out) \
        if not os.path.isabs(out) else out
    build(out, seed, a.project, seeded_from)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
